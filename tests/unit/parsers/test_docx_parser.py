"""DOCX 标题、段落和表格原始顺序测试。"""

from io import BytesIO

import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from PIL import Image

from app.models.documents import DocumentRef
from app.parsers.assets import AssetStore
from app.parsers.docx_parser import DocxParser
from app.parsers.render import render_document


def test_docx_keeps_section_content_and_tables(tmp_path):
    path = tmp_path / "sample.docx"
    document = Document()
    document.add_heading("技术要求", level=1)
    document.add_paragraph("像素不低于 640x480")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "参数"
    table.cell(0, 1).text = "要求"
    table.cell(1, 0).text = "精度"
    table.cell(1, 1).text = "±2℃"
    document.save(path)

    result = DocxParser().parse(
        DocumentRef(
            file_id="file-docx",
            file_name=path.name,
            source_path=path.name,
            data_role="rule_source",
            project_id="project-1",
        ),
        path,
    )

    assert result.parse_status == "parsed"
    assert result.sections[0].title == "技术要求"
    assert "640x480" in result.sections[0].content
    assert "±2℃" in result.sections[0].tables[0]


def image_node(paragraph, kind="vml"):
    container = OxmlElement("w:pict" if kind == "vml" else "w:drawing")
    tag = "{urn:schemas-microsoft-com:vml}imagedata" if kind == "vml" else qn("a:blip")
    node = etree.SubElement(container, tag)
    paragraph.add_run()._r.append(container)
    return node


def parse_fixture(document, tmp_path, **asset_options):
    path = tmp_path / "images.docx"
    document.save(path)
    assets = AssetStore(tmp_path / "file-docx", **asset_options)
    parsed = DocxParser().parse(
        DocumentRef(
            file_id="file-docx",
            file_name=path.name,
            source_path=path.name,
            data_role="evidence_material",
            bidder_id="bidder-test",
        ),
        path,
        assets=assets,
    )
    return parsed, assets, render_document(parsed)


@pytest.mark.parametrize("kind", ["vml", "drawingml"])
@pytest.mark.parametrize("in_table", [False, True])
def test_empty_image_node_has_no_readable_path(tmp_path, kind, in_table):
    doc = Document()
    paragraph = (
        doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
        if in_table
        else doc.add_paragraph()
    )
    paragraph.add_run("前文")
    node = image_node(paragraph, kind)
    node.set("{urn:schemas-microsoft-com:office:office}title", "")
    paragraph.add_run("后文")
    doc.add_paragraph("后续正常正文")

    parsed, assets, markdown = parse_fixture(doc, tmp_path)
    assert parsed.parse_status == "partial"
    assert len(parsed.warnings) == 1
    assert "空图片节点：未关联图片资源" in parsed.warnings[0]
    assert "没有可查看的图片路径" in parsed.warnings[0]
    assert "ValueError" not in markdown and "/images/" not in markdown
    assert (
        markdown.index("前文") < markdown.index("空图片节点") < markdown.index("后文")
    )
    assert "后续正常正文" in markdown
    assert not any(b.images for s in parsed.sections for b in s.blocks)
    assert not assets.root.exists()
    if in_table:
        assert len(parsed.sections[0].tables) == 1
        assert "空图片节点" in parsed.sections[0].tables[0]


@pytest.mark.parametrize(
    "kind", ["missing", "external_id", "external_link", "external_src", "non_image"]
)
def test_image_reference_errors_are_distinguished(tmp_path, kind):
    doc = Document()
    node = image_node(doc.add_paragraph("图片位置"))
    if kind == "missing":
        node.set(qn("r:id"), "rIdMissing")
        expected = "图片引用关系缺失"
    elif kind == "external_src":
        node.set("src", "https://example.invalid/remote.png")
        expected = "图片仅提供外部引用"
    elif kind == "non_image":
        node.set(
            qn("r:id"),
            doc.part.relate_to(doc.part.part_related_by(RT.STYLES), RT.STYLES),
        )
        expected = "图片引用指向非图片资源"
    else:
        rid = doc.part.relate_to(
            "https://example.invalid/remote.png", RT.IMAGE, is_external=True
        )
        node.set(qn("r:link") if kind == "external_link" else qn("r:id"), rid)
        expected = (
            "图片仅提供外部引用" if kind == "external_link" else "图片引用指向外部资源"
        )
    parsed, assets, markdown = parse_fixture(doc, tmp_path)
    assert expected in markdown and expected in parsed.warnings[0]
    assert "空图片节点" not in markdown
    assert "没有可查看的图片路径" in markdown
    assert not assets.root.exists()
    # 不把外部链接当作模型可用的图片地址，也不下载它。
    assert "https://example.invalid" not in markdown


@pytest.mark.parametrize(
    "payload,expected",
    [
        (b"", "已关联的图片资源内容为空"),
        (b"not-an-image", "已关联图片资源，但格式无法识别或数据已损坏"),
    ],
)
def test_invalid_embedded_image_data_keeps_reason(tmp_path, payload, expected):
    doc = Document()
    node = image_node(doc.add_paragraph("图片前文"))
    part = Part(PackURI("/word/media/broken.png"), "image/png", payload)
    node.set(qn("r:id"), doc.part.relate_to(part, RT.IMAGE))
    doc.add_paragraph("图片后文")
    parsed, assets, markdown = parse_fixture(doc, tmp_path)
    assert expected in parsed.warnings[0] and expected in markdown
    assert "空图片节点" not in markdown
    assert "图片前文" in markdown and "图片后文" in markdown
    assert not assets.root.exists()


@pytest.mark.parametrize(
    "limit,expected",
    [
        ({"max_pixels": 10}, "图片像素数超限"),
        ({"max_bytes": 2}, "图片字节数超限"),
    ],
)
def test_image_limit_error_keeps_specific_cause(tmp_path, limit, expected):
    doc = Document()
    data = BytesIO()
    Image.new("RGB", (20, 20), "white").save(data, format="PNG")
    doc.add_picture(BytesIO(data.getvalue()))
    parsed, _, markdown = parse_fixture(doc, tmp_path, **limit)
    assert expected in parsed.warnings[0] and expected in markdown
    assert "空图片节点" not in markdown
    assert not any(b.images for s in parsed.sections for b in s.blocks)


@pytest.mark.parametrize(
    "error,expected",
    [
        (OSError("decoder not available"), "decoder not available"),
        (
            PermissionError(13, "Permission denied", "/private/internal/image.png"),
            "Permission denied",
        ),
    ],
)
def test_decode_or_storage_failure_keeps_cause_without_server_path(
    tmp_path, monkeypatch, error, expected
):
    doc = Document()
    data = BytesIO()
    Image.new("RGB", (20, 20), "white").save(data, format="PNG")
    doc.add_picture(BytesIO(data.getvalue()))

    def fail(*_, **__):
        raise error

    monkeypatch.setattr(AssetStore, "save_image", fail)
    parsed, _, markdown = parse_fixture(doc, tmp_path)
    assert "解码或保存失败" in markdown and expected in parsed.warnings[0]
    assert "/private/internal" not in markdown
    assert "空图片节点" not in markdown
    assert "没有可查看的图片路径" in markdown


@pytest.mark.parametrize("attribute", ["r:id", "o:relid"])
def test_valid_vml_and_drawingml_survive_empty_node(tmp_path, attribute):
    doc = Document()
    data = BytesIO()
    Image.new("RGB", (20, 20), "white").save(data, format="PNG")
    doc.add_picture(BytesIO(data.getvalue()))
    rid, _ = doc.part.get_or_add_image(BytesIO(data.getvalue()))
    image_node(doc.add_paragraph("空节点所在段落"))
    node = image_node(doc.add_paragraph("正常旧式图片"))
    key = (
        qn("r:id")
        if attribute == "r:id"
        else "{urn:schemas-microsoft-com:office:office}relid"
    )
    node.set(key, rid)
    parsed, assets, markdown = parse_fixture(doc, tmp_path)
    images = [
        image
        for section in parsed.sections
        for block in section.blocks
        for image in block.images
    ]
    assert len(images) == 2 and len(parsed.warnings) == 1
    assert all(image.path == "file-docx/images/1.png" for image in images)
    assert markdown.count("![图片](file-docx/images/1.png)") == 2
    assert assets.read_image(images[0]) == assets.read_image(images[1])
