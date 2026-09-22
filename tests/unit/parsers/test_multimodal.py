"""Task19 结构、资产、PDF 与视觉调用边界；测试不访问模型服务。"""

import json

import pymupdf
import pytest
from docx import Document
from langchain_core.messages import AIMessage
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from PIL import Image

from app.models.documents import DocumentRef, ParsedDocument, ParsedImage, ParsedSection
from app.parsers import common, pdf_parser, service, vision, xlsx_parser
from app.parsers.assets import AssetStore, build_image_content, file_asset_root
from app.parsers.docx_parser import DocxParser
from app.parsers.image_messages import build_view_image_result
from app.parsers.image_parser import ImageParser
from app.parsers.pdf_parser import PdfParser
from app.parsers.render import render_document
from app.parsers.table_boundary import apply_boundary
from app.parsers.text_parser import TextParser
from app.parsers.xlsx_parser import XlsxParser


def reference(path, file_id="file-42"):
    return DocumentRef(
        file_id=file_id,
        file_name=path.name,
        source_path=path.name,
        data_role="rule_source",
        project_id="project-1",
    )


@pytest.fixture
def setup_parser(tmp_path, temp_settings, monkeypatch):
    for module in [common, service, vision, pdf_parser, xlsx_parser]:
        monkeypatch.setattr(module, "get_settings", lambda: temp_settings)
    assets = AssetStore(file_asset_root(temp_settings.storage.data_root, "file-42"))
    path = tmp_path / "picture.png"
    Image.new("RGB", (160, 100), "white").save(path)
    return assets, path, temp_settings


def test_docx_order_merged_cells_and_cell_images(setup_parser, tmp_path):
    assets, image, _ = setup_parser
    path = tmp_path / "mixed.docx"
    doc = Document()
    doc.add_heading("资料", 1)
    paragraph = doc.add_paragraph("前文")
    paragraph.add_run().add_picture(str(image))
    paragraph.add_run("后文")
    table = doc.add_table(rows=3, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "合并表头"
    table.cell(1, 0).merge(table.cell(2, 0)).text = "跨行"
    table.cell(1, 1).paragraphs[0].add_run().add_picture(str(image))
    table.cell(2, 2).text = "10万元"
    doc.add_paragraph("表后说明")
    doc.save(path)
    parsed = DocxParser().parse(reference(path), path, assets=assets)
    blocks = parsed.sections[0].blocks
    assert [b.kind for b in blocks] == ["text", "image", "text", "table", "text"]
    assert 'colspan="2"' in blocks[3].content and 'rowspan="2"' in blocks[3].content
    assert blocks[3].content.count("跨行") == 1
    assert len(blocks[3].images) == 1
    assert blocks[3].images[0].path in blocks[3].content
    assert parsed.parse_status == "partial"
    md = render_document(parsed)
    assert (
        md.index("前文")
        < md.index("后文")
        < md.index("合并表头")
        < md.index("表后说明")
    )
    assert "data:image" not in parsed.model_dump_json()


def test_xlsx_images_merge_and_empty_fields(setup_parser, tmp_path):
    assets, image, _ = setup_parser
    path = tmp_path / "table.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append(["A", "B", "C"])
    sheet.append(["value", None, "=1+1"])
    sheet.merge_cells("A3:B4")
    sheet["A3"] = "合并"
    sheet.add_image(ExcelImage(image), "C4")
    book.save(path)
    result = XlsxParser().parse(reference(path), path, assets=assets)
    table, picture = result.sections[0].blocks
    assert 'rowspan="2" colspan="2"' in table.content
    assert "[公式，未计算] =1+1" in table.content
    assert 'rowspan="1" colspan="1"></td>' in table.content
    assert "行 4 列 3" in picture.images[0].description
    assert result.parse_status == "partial"


def test_markdown_table_is_a_whole_structure(tmp_path):
    path = tmp_path / "text.md"
    path.write_text(
        "# 标题\n引导语\n\n| A | B |\n| --- | --- |\n| 1 | 2 |\n\n后文",
        encoding="utf-8",
    )
    parsed = TextParser().parse(reference(path), path)
    assert [b.kind for b in parsed.sections[0].blocks] == ["text", "table", "text"]
    assert "| 1 | 2 |" in parsed.sections[0].blocks[1].content


def test_legacy_json_keeps_version_and_missing_assets(tmp_path):
    old = ParsedDocument(
        document=reference(tmp_path / "old.docx"),
        parse_status="parsed",
        sections=[
            ParsedSection(
                section_id="s-1",
                title="旧标题",
                content="正文",
                tables=["旧表格"],
                image_count=1,
            )
        ],
    )
    loaded = ParsedDocument.model_validate_json(old.model_dump_json())
    assert loaded.schema_version == 1 and not loaded.sections[0].blocks
    assert "旧解析结果" in render_document(loaded)


def test_asset_security_and_tool_payload(setup_parser, tmp_path):
    assets, path, _ = setup_parser
    parsed = ImageParser().parse(reference(path), path, assets=assets)
    image = parsed.sections[0].blocks[0].images[0]
    result = build_view_image_result(
        parsed=parsed,
        image_path=image.path,
        assets=assets,
        allowed_file_ids={"file-42"},
        tool_call_id="call-1",
    )
    assert result.tool_call_id == "call-1"
    assert result.content[1]["image_url"]["url"].startswith("data:image/png;base64,")
    from langchain_openai.chat_models.base import _convert_message_to_dict

    wire = _convert_message_to_dict(result)
    assert wire["tool_call_id"] == "call-1" and wire["role"] == "tool"
    assert wire["content"][1]["type"] == "image_url"
    with pytest.raises(ValueError):
        build_view_image_result(
            parsed=parsed,
            image_path=image.path,
            assets=assets,
            allowed_file_ids={"file-other"},
            tool_call_id="call-1",
        )
    with pytest.raises(ValueError):
        build_image_content(parsed, "../../.env", assets)
    with pytest.raises(ValueError):
        file_asset_root(tmp_path, "../../private")
    with pytest.raises(ValueError):
        ParsedImage(
            path="images/../../secret.png", mime_type="image/png", width=1, height=1
        )
    outside = tmp_path / "outside.png"
    outside.write_bytes(path.read_bytes())
    target = assets.root / "images" / "link.png"
    target.symlink_to(outside)
    with pytest.raises(ValueError):
        assets.read_image(image.model_copy(update={"path": "images/link.png"}))
    with pytest.raises(ValueError):
        build_image_content(parsed, image.path, AssetStore(tmp_path / "file-other"))


def test_asset_limits_and_idempotent_content(setup_parser):
    assets, path, _ = setup_parser
    one = assets.save_image(path.read_bytes())
    two = assets.save_image(path.read_bytes())
    assert one.path == two.path
    with pytest.raises(ValueError):
        AssetStore(assets.root, max_pixels=10).save_image(path.read_bytes())
    with pytest.raises(ValueError):
        AssetStore(assets.root, max_bytes=2).save_image(path.read_bytes())


def make_pdf(path, *, scan=None, table=False, pages=1):
    pdf = pymupdf.open()
    for index in range(pages):
        page = pdf.new_page(width=300, height=400)
        if scan:
            page.insert_image(pymupdf.Rect(20, 80, 260, 250), filename=str(scan))
        if table:
            for y in [60, 110, 160]:
                page.draw_line((20, y), (280, y))
            for x in [20, 150, 280]:
                page.draw_line((x, 60), (x, 160))
            for xy, text in [
                ((25, 85), "Field"),
                ((155, 85), "Value"),
                ((25, 135), f"Row{index}"),
                ((155, 135), str(index + 10)),
            ]:
                page.insert_text(xy, text)
        if not scan:
            page.insert_text((20, 30), "Native text")
    pdf.save(path)
    pdf.close()


def test_pdf_native_table_and_page_locations(setup_parser, tmp_path):
    assets, _, _ = setup_parser
    path = tmp_path / "native.pdf"
    make_pdf(path, table=True)
    parsed = PdfParser().parse(reference(path), path, assets=assets)
    blocks = parsed.sections[0].blocks
    assert any("Native text" in b.content for b in blocks)
    table = next(b for b in blocks if b.kind == "table")
    assert "Row0" in table.content and "10" in table.content
    assert table.page_sources[0].page_number == 1
    assert table.page_sources[0].image_path == table.images[0].path
    assert table.images[0].purpose == "reference"
    assert parsed.parse_status == "parsed"


def test_scan_pdf_is_not_empty_or_falsely_readable(setup_parser, tmp_path):
    assets, image, _ = setup_parser
    path = tmp_path / "scan.pdf"
    make_pdf(path, scan=image)
    result = PdfParser().parse(reference(path), path, assets=assets)
    assert result.parse_status == "ocr_required"
    assert any(b.images for b in result.sections[0].blocks)


class FakeModel:
    def __init__(self, values):
        self.values = iter(values)
        self.calls = []

    async def ainvoke(self, messages):
        self.calls.append(messages)
        value = next(self.values)
        if isinstance(value, Exception):
            raise value
        if isinstance(value, AIMessage):
            return value
        return AIMessage(
            content=json.dumps(value, ensure_ascii=False),
            response_metadata={"finish_reason": "stop"},
        )


@pytest.mark.asyncio
async def test_vision_cached_and_no_base64_persisted(setup_parser):
    assets, path, _ = setup_parser
    native = ImageParser().parse(reference(path), path, assets=assets)
    model = FakeModel(
        [{"status": "readable", "content": "编号 123；金额 10万元", "reason": ""}]
    )
    one = await vision.enrich_document(native, assets, model=model)
    two = await vision.enrich_document(native, assets, model=FakeModel([]))
    assert one.parse_status == two.parse_status == "parsed"
    assert one.sections[0].content == two.sections[0].content
    assert len(model.calls) == 1
    assert model.calls[0][1].content[1]["type"] == "image_url"
    for file in assets.root.rglob("*.json"):
        assert "base64" not in file.read_text("utf-8")


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "reply",
    [
        RuntimeError("offline"),
        AIMessage(
            content='{"status":"readable"}',
            response_metadata={"finish_reason": "length"},
        ),
        {"status": "readable", "content": "<script>bad()</script>", "reason": ""},
    ],
)
async def test_vision_failure_keeps_original_and_can_retry(setup_parser, reply):
    assets, path, _ = setup_parser
    native = ImageParser().parse(reference(path), path, assets=assets)
    failed = await vision.enrich_document(native, assets, model=FakeModel([reply]))
    image = failed.sections[0].blocks[0].images[0]
    assert image.read_status == "unreadable" and image.transcription == ""
    assert (assets.root / image.path).exists()
    retried = await vision.enrich_document(
        failed,
        assets,
        model=FakeModel([{"status": "readable", "content": "成功", "reason": ""}]),
    )
    assert retried.parse_status == "parsed"


@pytest.mark.asyncio
async def test_missing_vision_key_not_replaced_with_agent_key(
    setup_parser, monkeypatch
):
    assets, path, settings = setup_parser
    monkeypatch.setattr(
        vision,
        "get_settings",
        lambda: settings.model_copy(update={"parser_vision_model": None}),
    )
    native = ImageParser().parse(reference(path), path, assets=assets)
    result = await vision.enrich_document(native, assets)
    assert result.parse_status == "ocr_required"
    assert "配置缺失" in result.warnings[-1]


@pytest.mark.asyncio
async def test_pdf_cross_page_merge_and_preserved_sources(setup_parser, tmp_path):
    assets, _, _ = setup_parser
    path = tmp_path / "tables.pdf"
    make_pdf(path, table=True, pages=2)
    native = PdfParser().parse(reference(path), path, assets=assets)
    model = FakeModel(
        [
            {
                "action": "merge",
                "drop_header_rows": 1,
                "join_boundary_row": False,
                "reason": "续表",
            }
        ]
    )
    result = await vision.enrich_document(native, assets, model=model)
    tables = [b for s in result.sections for b in s.blocks if b.kind == "table"]
    assert len(tables) == 1
    assert [p.page_number for p in tables[0].page_sources] == [1, 2]
    assert len(tables[0].images) == 2 and tables[0].content_origin == "vision_assisted"
    assert len([b for s in native.sections for b in s.blocks if b.kind == "table"]) == 2


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "decision",
    [
        {
            "action": "separate",
            "drop_header_rows": 0,
            "join_boundary_row": False,
            "reason": "不同表",
        },
        {
            "action": "uncertain",
            "drop_header_rows": 0,
            "join_boundary_row": False,
            "reason": "表题不清",
        },
        {
            "action": "merge",
            "drop_header_rows": 2,
            "join_boundary_row": False,
            "reason": "续表",
        },
    ],
)
async def test_nonmerge_or_invalid_merge_never_drops_second_page(
    setup_parser, tmp_path, decision
):
    assets, _, _ = setup_parser
    path = tmp_path / "tables.pdf"
    make_pdf(path, table=True, pages=2)
    native = PdfParser().parse(reference(path), path, assets=assets)
    result = await vision.enrich_document(native, assets, model=FakeModel([decision]))
    assert len([b for s in result.sections for b in s.blocks if b.kind == "table"]) == 2
    assert "Row1" in result.sections[1].tables[0]


class FakeRepository:
    def update_parse_status(self, file_id, status):
        self.status = status


def test_bundle_atomic_pointer_and_old_images_preserved(setup_parser, monkeypatch):
    assets, path, settings = setup_parser
    native = ImageParser().parse(reference(path), path, assets=assets)
    repo = FakeRepository()
    original = service.persist_parse_result(native, repo)
    entry = settings.storage.data_root / "parsed/file-42.json"
    previous = entry.read_bytes()
    real_write = service.atomic_write

    def broken(target, content):
        if target == entry:
            raise OSError("disk failure")
        real_write(target, content)

    monkeypatch.setattr(service, "atomic_write", broken)
    changed = native.model_copy(deep=True)
    changed.warnings.append("新版本")
    with pytest.raises(OSError):
        service.persist_parse_result(changed, repo)
    assert entry.read_bytes() == previous
    assert (settings.storage.data_root / "parsed" / original.markdown_path).exists()
    assert assets.read_image(native.sections[0].blocks[0].images[0])


@pytest.mark.asyncio
async def test_public_prepare_entry_does_not_call_vision_when_disabled(setup_parser):
    assets, path, _ = setup_parser

    class Store:
        def resolve(self, file_id):
            return path

    model = FakeModel([])
    result = await service.prepare_document(
        reference(path),
        file_store=Store(),
        repository=FakeRepository(),
        vision_enabled=False,
        vision_model=model,
    )
    assert not model.calls
    assert result.schema_version == 2 and result.parse_status == "ocr_required"
    assert list((assets.root / "checkpoints").glob("*.json"))


def test_merge_validation_keeps_duplicate_rows_and_empty_values():
    left = (
        "<table><tr><td>A</td><td>B</td></tr><tr><td>value</td><td></td></tr></table>"
    )
    right = (
        "<table><tr><td>A</td><td>B</td></tr><tr><td>value</td><td></td></tr></table>"
    )
    with pytest.raises(ValueError):
        apply_boundary(left, right, drop_header_rows=2, join_boundary_row=False)
    result = apply_boundary(left, right, drop_header_rows=1, join_boundary_row=False)
    assert result.count("value") == 2


@pytest.mark.asyncio
async def test_prepare_entry_with_mock_vision_publishes_transcription(setup_parser):
    assets, path, _ = setup_parser

    class Store:
        def resolve(self, file_id):
            return path

    model = FakeModel(
        [{"status": "readable", "content": "主体名称：测试公司", "reason": ""}]
    )
    result = await service.prepare_document(
        reference(path),
        file_store=Store(),
        repository=FakeRepository(),
        vision_enabled=True,
        vision_model=model,
    )
    assert result.parse_status == "parsed"
    assert "测试公司" in service.load_parsed_document("file-42").sections[0].content
    assert "测试公司" in (assets.root / "document.md").read_text("utf-8")


def test_pdf_rotation_sources_in_display_coordinates(setup_parser, tmp_path):
    assets, _, _ = setup_parser
    path = tmp_path / "rotated.pdf"
    with pymupdf.open() as pdf:
        page = pdf.new_page(width=300, height=400)
        page.insert_text((20, 30), "Rotation")
        page.set_rotation(90)
        pdf.save(path)
    result = PdfParser().parse(reference(path), path, assets=assets)
    source = result.sections[0].blocks[0].page_sources[0]
    assert 0 <= source.bbox[0] < source.bbox[2] <= 400
    assert 0 <= source.bbox[1] < source.bbox[3] <= 300
    with pymupdf.open(path) as unchanged:
        assert unchanged[0].rotation == 90


def test_v2_empty_blocks_do_not_resurrect_legacy_table(tmp_path):
    parsed = ParsedDocument(
        schema_version=2,
        document=reference(tmp_path / "table.pdf"),
        parse_status="parsed",
        sections=[
            ParsedSection(section_id="page-2", title="第2页", tables=["已被合并的旧表"])
        ],
    )
    assert parsed.sections[0].tables == []
    assert "旧表" not in render_document(parsed)
    assert (
        ParsedDocument.model_validate_json(parsed.model_dump_json()).sections[0].tables
        == []
    )


def test_short_assets_reuse_and_concurrent_allocation(setup_parser):
    import io
    from concurrent.futures import ThreadPoolExecutor

    assets, path, _ = setup_parser
    first = assets.save_image(path.read_bytes())
    assert first.path == "images/1.png"
    payloads = []
    for color in ["red", "green", "blue"]:
        buffer = io.BytesIO()
        Image.new("RGB", (20, 20), color).save(buffer, format="PNG")
        payloads.append(buffer.getvalue())
    with ThreadPoolExecutor(max_workers=6) as pool:
        images = list(pool.map(assets.save_image, payloads * 2))
    assert {i.path for i in images} == {"images/2.png", "images/3.png", "images/4.png"}
    assert assets.save_image(path.read_bytes()).path == first.path
    assert assets.read_image(first)


@pytest.mark.asyncio
async def test_four_page_merge_published_without_duplicate_tables(
    setup_parser, tmp_path
):
    assets, _, settings = setup_parser
    path = tmp_path / "four.pdf"
    make_pdf(path, table=True, pages=4)
    native = PdfParser().parse(reference(path), path, assets=assets)
    # 重现原缺陷：被合并页只剩空 blocks，但原生阶段曾有 tables 投影。
    for section in native.sections:
        section.blocks = [b for b in section.blocks if b.kind == "table"]
        section.project_blocks()
    model = FakeModel(
        [
            {
                "action": "merge",
                "drop_header_rows": 1,
                "join_boundary_row": False,
                "reason": "相邻续表",
            }
        ]
        * 3
    )
    result = await vision.enrich_document(native, assets, model=model)
    stored = service.persist_parse_result(result, FakeRepository())
    assert len(model.calls) == 3
    assert len(stored.sections[0].blocks[0].page_sources) == 4
    assert all(not s.tables and not s.blocks for s in stored.sections[1:])
    markdown = (
        settings.storage.data_root / "parsed" / stored.markdown_path
    ).read_text()
    for index in range(4):
        assert markdown.count(f"Row{index}") == 1
    for call in model.calls:
        assert sum(v["type"] == "image_url" for v in call[1].content) == 2
    assert "![图片]" not in markdown  # 回查图只作链接，不再当正文图片展开。


def test_split_row_preserves_cell_binding_and_rejects_unsafe_merge():
    from app.parsers.table_boundary import (
        BoundaryError,
        rows_html,
        table_rows,
        text_value,
    )

    left = rows_html([["编号", "名称", "说明"], ["12", "甲", "上半"]])
    right = rows_html(
        [["编号", "名称", "说明"], ["", "", "下半"], ["13", "乙", "完整"]]
    )
    result = apply_boundary(left, right, drop_header_rows=1, join_boundary_row=True)
    assert [[text_value(c) for c in row] for row in table_rows(result)] == [
        ["编号", "名称", "说明"],
        ["12", "甲", "上半下半"],
        ["13", "乙", "完整"],
    ]
    with pytest.raises(BoundaryError, match="首列"):
        apply_boundary(
            left,
            rows_html([["编号", "名称", "说明"], ["13", "乙", "完整"]]),
            drop_header_rows=1,
            join_boundary_row=True,
        )
    with pytest.raises(BoundaryError, match="跨行"):
        apply_boundary(
            '<table><tr><td colspan="2">标题</td></tr></table>',
            left,
            drop_header_rows=0,
            join_boundary_row=False,
        )


@pytest.mark.asyncio
async def test_boundary_diagnostic_has_scope_without_sensitive_error(
    setup_parser, tmp_path
):
    assets, _, _ = setup_parser
    path = tmp_path / "failure.pdf"
    make_pdf(path, table=True, pages=2)
    native = PdfParser().parse(reference(path), path, assets=assets)
    result = await vision.enrich_document(
        native, assets, model=FakeModel([RuntimeError("secret-key-and-content")])
    )
    assert sum(len(s.tables) for s in result.sections) == 2
    diagnostic = next((assets.root / "checkpoints/diagnostics").glob("*.json"))
    data = json.loads(diagnostic.read_text())
    assert data["pages"] == [1, 2] and data["stage"] == "request"
    assert "secret-key" not in diagnostic.read_text() + result.model_dump_json()


@pytest.mark.asyncio
async def test_image_context_cache_changes_with_source_metadata(setup_parser):
    assets, path, _ = setup_parser
    native = ImageParser().parse(reference(path), path, assets=assets)
    reply = {"status": "readable", "content": "原图内容", "reason": ""}
    model = FakeModel([reply, reply])
    await vision.enrich_document(native, assets, model=model)
    changed = native.model_copy(deep=True)
    changed.sections[0].title = "另一处来源"
    await vision.enrich_document(changed, assets, model=model)
    assert len(model.calls) == 2
    context = json.loads(model.calls[1][1].content[0]["text"])
    assert context["来源定位"]["章节或工作表"] == "另一处来源"
    assert context["来源定位"]["页码"] == []  # 没有真实分页不能编造页码。


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "configured,override,expected",
    [
        (False, None, False),
        (True, None, True),
        (True, False, False),
        (False, True, True),
    ],
)
async def test_prepare_vision_config_and_explicit_override(
    setup_parser, configured, override, expected
):
    _, path, settings = setup_parser
    settings.parsing.vision_enabled = configured

    class Store:
        def resolve(self, file_id):
            return path

    model = FakeModel([{"status": "readable", "content": "测试内容", "reason": ""}])
    await service.prepare_document(
        reference(path),
        file_store=Store(),
        repository=FakeRepository(),
        vision_enabled=override,
        vision_model=model,
    )
    assert bool(model.calls) == expected


def test_pdf_image_adaptive_resolution_keeps_page_overlay(setup_parser, tmp_path):
    assets, _, _ = setup_parser
    image = tmp_path / "large.png"
    Image.new("RGB", (1000, 500), "white").save(image)
    path = tmp_path / "quality.pdf"
    with pymupdf.open() as pdf:
        page = pdf.new_page(width=400, height=300)
        page.insert_text((20, 20), "Native heading")
        page.insert_image(pymupdf.Rect(20, 50, 260, 170), filename=str(image))
        page.draw_rect(pymupdf.Rect(20, 50, 140, 170), color=(1, 0, 0), fill=(1, 0, 0))
        pdf.save(path)
    before = path.read_bytes()
    parsed = PdfParser().parse(reference(path), path, assets=assets)
    picture = next(
        i
        for s in parsed.sections
        for b in s.blocks
        for i in b.images
        if i.purpose == "content"
    )
    assert picture.width >= 1000 and picture.height >= 500
    with Image.open(assets.root / picture.path) as result:
        assert result.getpixel((50, 50)) == (255, 0, 0)
    assert path.read_bytes() == before


def test_legacy_hash_asset_readable_and_failed_registration_never_overwrites(
    setup_parser, monkeypatch
):
    import hashlib

    import app.parsers.assets as asset_module

    assets, path, _ = setup_parser
    first = assets.save_image(path.read_bytes())
    payload = assets.read_image(first)
    old_path = f"images/{hashlib.sha256(payload).hexdigest()}.png"
    (assets.root / old_path).write_bytes(payload)
    assert assets.read_image(first.model_copy(update={"path": old_path})) == payload
    real_write = asset_module.atomic_write

    def fail_image_write(target, data):
        if target.name == "2.png":
            raise OSError("simulated interrupted publish")
        real_write(target, data)

    monkeypatch.setattr(asset_module, "atomic_write", fail_image_write)
    changed = path.with_name("changed.png")
    Image.new("RGB", (25, 25), "black").save(changed)
    with pytest.raises(OSError):
        assets.save_image(changed.read_bytes())
    assert assets.read_image(first) == payload
    monkeypatch.setattr(asset_module, "atomic_write", real_write)
    assert assets.save_image(changed.read_bytes()).path == "images/3.png"


@pytest.mark.parametrize(
    "flags,expected", [([], None), (["--vision"], True), (["--no-vision"], False)]
)
def test_cli_vision_switches(monkeypatch, flags, expected):
    import sys
    from types import SimpleNamespace

    from app.parsers import __main__ as cli

    document = SimpleNamespace(file_id="file-42")
    monkeypatch.setattr(
        cli, "FileRepository", lambda: SimpleNamespace(require=lambda _: document)
    )
    monkeypatch.setattr(cli, "document_ref_from_record", lambda value: value)
    captured = []

    async def prepare(ref, *, vision_enabled):
        captured.append(vision_enabled)
        return SimpleNamespace(
            parse_status="parsed", markdown_path="test.md", warnings=[]
        )

    monkeypatch.setattr(cli, "prepare_document", prepare)
    monkeypatch.setattr(sys, "argv", ["parsers", "file-42", *flags])
    cli.main()
    assert captured == [expected]


@pytest.mark.asyncio
async def test_vision_enabled_but_no_content_makes_no_request(setup_parser, tmp_path):
    assets, _, _ = setup_parser
    path = tmp_path / "plain.txt"
    path.write_text("普通原文", encoding="utf-8")
    parsed = TextParser().parse(reference(path), path)
    model = FakeModel([])
    await vision.enrich_document(parsed, assets, model=model)
    assert not model.calls


@pytest.mark.asyncio
async def test_uncertain_boundary_is_not_hidden_by_later_success(
    setup_parser, tmp_path
):
    assets, _, _ = setup_parser
    path = tmp_path / "uncertain.pdf"
    make_pdf(path, table=True, pages=4)
    parsed = PdfParser().parse(reference(path), path, assets=assets)
    merge = {
        "action": "merge",
        "drop_header_rows": 1,
        "join_boundary_row": False,
        "reason": "续表",
    }
    uncertain = {
        "action": "uncertain",
        "drop_header_rows": 0,
        "join_boundary_row": False,
        "reason": "无法确认关系",
    }
    result = await vision.enrich_document(
        parsed, assets, model=FakeModel([merge, uncertain, merge])
    )
    tables = [b for s in result.sections for b in s.blocks if b.kind == "table"]
    assert [[p.page_number for p in t.page_sources] for t in tables] == [[1, 2], [3, 4]]
    assert result.parse_status == "partial"
    assert "第 2→3 页" in render_document(result)


@pytest.mark.asyncio
async def test_diagnostics_distinguish_json_and_field_validation(setup_parser):
    assets, _, _ = setup_parser
    try:
        json.loads('{"content": invalid}')
    except json.JSONDecodeError as exc:
        await vision._diagnostic(assets, stage="image", pages=[1], exc=exc)
    from pydantic import ValidationError

    try:
        vision.ImageReading.model_validate({"content": "sensitive material"})
    except ValidationError as exc:
        await vision._diagnostic(assets, stage="image", pages=[1], exc=exc)
    records = [
        json.loads(p.read_text())
        for p in (assets.root / "checkpoints/diagnostics").glob("*.json")
    ]
    assert any("json_position" in record for record in records)
    assert any(record.get("validation") for record in records)
    assert "sensitive material" not in json.dumps(records)
