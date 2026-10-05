"""XLSX 多 Sheet 解析测试。"""

from openpyxl import Workbook
from openpyxl.styles import Font
from lxml import etree
from datetime import datetime, time

import pytest

from app.models.documents import DocumentRef
from app.parsers.xlsx_parser import XlsxParser


def test_each_sheet_becomes_a_section(tmp_path):
    path = tmp_path / "sample.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "包1"
    sheet.append(["物资", "数量"])
    sheet.append(["红外热成像仪", 2])
    workbook.create_sheet("包2").append(["物资", "数量"])
    workbook.save(path)

    result = XlsxParser().parse(
        DocumentRef(
            file_id="file-xlsx",
            file_name=path.name,
            source_path=path.name,
            data_role="rule_source",
            project_id="project-1",
        ),
        path,
    )

    assert result.parse_status == "parsed"
    assert [section.title for section in result.sections] == [
        "Sheet: 包1",
        "Sheet: 包2",
    ]
    assert "红外热成像仪" in result.sections[0].tables[0]


def parse_book(tmp_path, book):
    path = tmp_path / "formatted.xlsx"
    book.save(path)
    result = XlsxParser().parse(DocumentRef(
        file_id="file-xlsx", file_name=path.name, source_path=path.name,
        data_role="rule_source", project_id="project-1",
    ), path)
    return etree.HTML(result.sections[0].blocks[0].content)


def test_worksheet_keeps_layout_metadata_and_original_grid(tmp_path):
    book = Workbook()
    sh = book.active
    sh["A2"] = "复检页明细"
    sh["A3"], sh["B3"], sh["C3"] = "供应商", "演练包", "说明"
    sh.append([])
    for c, title in zip(sh[6], ["序号", "日期", "说明"]):
        c.value, c.font = title, Font(bold=True)
    sh["A7"], sh["B7"], sh["C7"] = 1, datetime(2026, 9, 26), "复检"
    sh["B7"].number_format = "yyyy-mm-dd"
    sh["A9"] = "末行"
    sh.column_dimensions["A"].width = 8
    sh.column_dimensions["B"].width = 22
    root = parse_book(tmp_path, book)
    rows = root.xpath("//table/tr")
    assert len(rows) == 9  # 解析层不删除留白、内部空记录或改变原始行号。
    assert rows[0].get("data-source-row") == "1"
    assert not "".join(rows[0].itertext())
    assert len(rows[5].xpath("./th")) == 3
    assert len(rows[2].xpath("./td")) == 3
    assert rows[6][1].text == "2026-09-26"
    assert rows[6][1].get("data-raw-value") == "2026-09-26 00:00:00"
    assert rows[6][1].get("data-number-format") == "yyyy-mm-dd"
    assert root.xpath("//col/@data-width")[:2] == ["8.00", "22.00"]


@pytest.mark.parametrize("value,fmt,expected", [
    (datetime(2026, 9, 26), "yyyy-mm-dd", "2026-09-26"),
    (datetime(2026, 9, 26, 14, 30, 45), "yyyy-mm-dd hh:mm", "2026-09-26 14:30"),
    (datetime(2026, 9, 6), 'yyyy"年"m"月"d"日"', "2026年9月6日"),
    (datetime(2026, 9, 6), '"Date "YYYY-MM-DD', "Date 2026-09-06"),
    (time(9, 5, 6), "hh:mm:ss", "09:05:06"),
    (time(9, 5, 6), "mm:ss", "05:06"),
    (datetime(2026, 9, 26), "dddd", "2026-09-26 00:00:00"),
    (201250, "#,##0", "201,250"),
    (0.125, "0.0%", "12.5%"),
    (0.125, "0.00", "0.13"),
    (-0.125, "0.00", "-0.13"),
    (71, "0000", "0071"),
    ("=1+1", "0.00", "[公式，未计算] =1+1"),
    ('<b>&"原值', "General", '<b>&"原值'),
    (12.34, '0.00;[Red]-0.00', "12.34"),
])
def test_display_formats_are_conservative_and_retain_raw_value(tmp_path, value, fmt, expected):
    book = Workbook()
    cell = book.active["A1"]
    cell.value, cell.number_format = value, fmt
    root = parse_book(tmp_path, book)
    parsed = root.xpath("//td")[0]
    assert parsed.text == expected
    if expected != str(value) and cell.data_type != "f":
        assert parsed.get("data-raw-value") == str(value)


def test_title_and_notes_are_not_mistaken_for_header(tmp_path):
    book = Workbook()
    sh = book.active
    sh.merge_cells("A1:C1")
    sh["A1"], sh["A1"].font = "人员安排", Font(bold=True)
    sh.append(["项目说明"])
    sh.append(["姓名", "岗位", "驻场"])
    for c in sh[3]:
        c.font = Font(bold=True)
    sh.append(["甲", "负责人", "是"])
    root = parse_book(tmp_path, book)
    assert len(root.xpath("//tr[3]/th")) == 3
    assert root.xpath("//tr[1]/td")[0].get("colspan") == "3"
    assert not root.xpath("//tr[1]/th | //tr[4]/th")


def test_later_bold_summary_is_not_header(tmp_path):
    book = Workbook()
    sh = book.active
    sh.append(["甲", 1])
    sh.append(["汇总", "说明"])
    for c in sh[2]:
        c.font = Font(bold=True)
    root = parse_book(tmp_path, book)
    assert not root.xpath("//th")
