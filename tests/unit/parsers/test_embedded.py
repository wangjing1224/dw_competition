"""自造无隐私嵌入文件，覆盖多层 Office/OLE/压缩包及失败隔离。"""

import struct
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZipFile

import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement
from lxml import etree
from openpyxl import Workbook
from PIL import Image

from app.core.config import get_settings
from app.models.documents import DocumentRef
from app.parsers.assets import AssetStore
from app.parsers.docx_parser import DocxParser
from app.parsers.embedded import EmbeddedParser, _native_payload
from app.parsers.render import render_document
from app.retrieval.chunker import chunk_parsed_document


def document_bytes(doc):
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def compound(stream_name, payload):
    """最小单流 CFB fixture：大流扇区+目录+FAT，不包含真实材料。"""
    payload = payload.ljust(4096, b"\0")
    sectors = (len(payload) + 511) // 512
    fat_count = 1
    while fat_count * 128 < sectors + 1 + fat_count:
        fat_count += 1
    assert fat_count <= 109
    free, end = 0xFFFFFFFF, 0xFFFFFFFE
    header = bytearray(512)
    header[:8] = bytes.fromhex("d0cf11e0a1b11ae1")
    struct.pack_into("<HHHHH", header, 24, 0x3E, 3, 0xFFFE, 9, 6)
    struct.pack_into("<IIIIIIIII", header, 40, 0, fat_count, sectors, 0, 4096, end, 0, end, 0)
    struct.pack_into("<109I", header, 76, *range(sectors + 1, sectors + 1 + fat_count), *([free] * (109 - fat_count)))

    def entry(name, kind, start, size, child=free):
        data = bytearray(128)
        label = name.encode("utf-16le") + b"\0\0"
        data[:len(label)] = label
        struct.pack_into("<HBBIII", data, 64, len(label), kind, 1, free, free, child)
        struct.pack_into("<IQ", data, 116, start, size)
        return data

    directory = entry("Root Entry", 5, end, 0, 1) + entry(stream_name, 2, 0, len(payload))
    fat = [*range(1, sectors), end, end, *([0xFFFFFFFD] * fat_count)]
    return bytes(header) + payload.ljust(sectors * 512, b"\0") + directory.ljust(512, b"\0") + struct.pack(f"<{128 * fat_count}I", *(fat + [free] * (128 * fat_count - len(fat))))


def ole_native(name, payload):
    body = (b"\x02\0" + name.encode("gb18030") + b"\0C:\\private\\input\0"
            + b"\0" * 8 + b"C:\\private\\temp\0" + struct.pack("<I", len(payload)) + payload)
    return compound("\x01Ole10Native", struct.pack("<I", len(body)) + body)


def embed(paragraph, data, *, rid=1, external=False, reltype=RT.OLE_OBJECT):
    obj = OxmlElement("w:object")
    # 不能读的预览图标应被跳过，不应遮盖真正可解析的附件。
    etree.SubElement(obj, "{urn:schemas-microsoft-com:vml}imagedata")
    node = etree.SubElement(obj, "{urn:schemas-microsoft-com:office:office}OLEObject")
    node.set("Type", "Embed")
    target = ("https://example.invalid/document" if external else Part(
        PackURI(f"/word/embeddings/oleObject{rid}.bin"), "application/octet-stream", data
    ))
    from docx.oxml.ns import qn
    node.set(qn("r:id"), paragraph.part.relate_to(target, reltype, is_external=external))
    paragraph.add_run()._r.append(obj)


def parse(doc, tmp_path):
    path = tmp_path / "parent.docx"
    doc.save(path)
    ref = DocumentRef(file_id="file-nested", file_name=path.name, source_path=path.name,
                      data_role="rule_source", project_id="project-test")
    assets = AssetStore(tmp_path / ref.file_id)
    result = DocxParser().parse(ref, path, assets=assets)
    return result, assets, render_document(result)


def workbook_bytes():
    book = Workbook()
    sheet = book.active
    sheet.title = "评分细则"
    sheet.merge_cells("A1:B1")
    sheet["A1"] = "评分条件"
    sheet["A2"], sheet["B2"] = "技术指标", 60
    book.create_sheet("说明")["A1"] = "不得借用其他标段的权重"
    out = BytesIO()
    book.save(out)
    return out.getvalue()


@pytest.mark.parametrize("in_table", [False, True])
def test_multilevel_docx_preserves_order_images_and_parent_scope(tmp_path, in_table):
    leaf = Document()
    leaf.add_heading("内层要求", 1)
    p = leaf.add_paragraph("评分公式 x")
    p.add_run("n").font.superscript = True
    leaf.add_table(rows=1, cols=1).cell(0, 0).text = "最内层完整评分表"
    picture = BytesIO()
    Image.new("RGB", (20, 20), "white").save(picture, format="PNG")
    leaf.add_picture(BytesIO(picture.getvalue()))
    middle = Document()
    p = middle.add_paragraph("中层前文")
    embed(p, compound("Package", document_bytes(leaf)))
    p.add_run("中层后文")
    parent = Document()
    p = parent.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0] if in_table else parent.add_paragraph()
    p.add_run("父文件前文")
    embed(p, compound("Package", document_bytes(middle)))
    p.add_run("父文件后文")
    parsed, assets, markdown = parse(parent, tmp_path)
    values = ["父文件前文", "中层前文", "评分公式 x^(n)", "最内层完整评分表", "中层后文", "父文件后文"]
    assert [markdown.index(v) for v in values] == sorted(markdown.index(v) for v in values)
    assert not parsed.warnings
    assert len({s.section_id for s in parsed.sections}) == len(parsed.sections)
    images = [i for s in parsed.sections for b in s.blocks for i in b.images]
    assert len(images) == 1 and images[0].path == "file-nested/images/1.png"
    assert assets.read_image(images[0])
    chunks = chunk_parsed_document(parsed, max_tokens=1000, overlap_tokens=20)
    assert all(c.file_id == "file-nested" and c.project_id == "project-test" for c in chunks)
    assert any("word/embeddings/" in (c.section + c.content) for c in chunks)


def test_native_zip_xlsx_keeps_all_sheets_and_continues_after_bad_member(tmp_path):
    out = BytesIO()
    with ZipFile(out, "w") as z:
        z.writestr("../outside.xlsx", workbook_bytes())
        z.writestr("评分/制造商.xlsx", workbook_bytes())
        z.writestr("坏附件.docx", b"not-a-document")
        z.writestr("评分/代理商.xlsx", workbook_bytes())
    doc = Document()
    embed(doc.add_paragraph("包前"), ole_native("评分.zip", out.getvalue()))
    doc.add_paragraph("包后")
    parsed, _, text = parse(doc, tmp_path)
    assert text.count("不得借用其他标段的权重") == 2
    assert text.count("技术指标") == 2
    assert 'colspan="2"' in text and "代理商.xlsx" in text
    assert "不安全的压缩包成员路径" in text and "暂不支持此附件格式" in text
    assert "C:\\private" not in text
    assert "包后" in text and parsed.parse_status == "partial"
    assert not (tmp_path.parent / "outside.xlsx").exists()


@pytest.mark.parametrize("mode", ["external", "missing", "wrong", "corrupt"])
def test_unreadable_objects_are_explicit_not_silent_or_fatal(tmp_path, mode):
    doc = Document()
    p = doc.add_paragraph("前文")
    embed(p, b"broken", external=mode == "external", reltype=RT.IMAGE if mode == "wrong" else RT.OLE_OBJECT)
    if mode == "missing":
        from docx.oxml.ns import qn
        p._p.find(".//{urn:schemas-microsoft-com:office:office}OLEObject").set(qn("r:id"), "absent")
    doc.add_paragraph("之后有效正文")
    parsed, _, text = parse(doc, tmp_path)
    assert parsed.parse_status == "partial" and len(parsed.warnings) == 1
    assert "未展开" in text and "之后有效正文" in text
    assert "空图片节点" not in text and "https://" not in text


@pytest.mark.parametrize("setting,value", [("max_embedded_depth", 1), ("max_embedded_files", 1), ("max_embedded_bytes", 32)])
def test_attachment_limits_do_not_drop_parent_text(tmp_path, monkeypatch, setting, value):
    monkeypatch.setattr(get_settings().parsing, setting, value)
    doc = Document()
    embed(doc.add_paragraph("前文"), compound("Package", workbook_bytes()))
    doc.add_paragraph("后文")
    parsed, _, text = parse(doc, tmp_path)
    assert parsed.parse_status == "partial" and "上限" in text and "后文" in text


def test_xls_uses_real_reader_and_preserves_merges_and_value_types(tmp_path, monkeypatch):
    import xlrd
    cell = lambda kind, value: SimpleNamespace(ctype=kind, value=value)
    cells = [[cell(1, "权重"), cell(0, ""), cell(1, "说明")],
             [cell(2, 60.0), cell(2, 0.35), cell(4, 1)]]
    sheet = SimpleNamespace(name="Sheet2", nrows=2, ncols=3, merged_cells=[(0, 1, 0, 2)],
                            cell=lambda r, c: cells[r][c])
    called = []
    def opened(**kwargs):
        assert kwargs["formatting_info"] and kwargs["file_contents"].startswith(bytes.fromhex("d0cf11e0"))
        return SimpleNamespace(sheets=lambda: [sheet], release_resources=lambda: called.append(True))
    monkeypatch.setattr(xlrd, "open_workbook", opened)
    doc = Document()
    embed(doc.add_paragraph(), compound("Workbook", b"mock BIFF"))
    parsed, _, text = parse(doc, tmp_path)
    assert 'colspan="2"' in text and ">60<" in text and "0.35" in text and "TRUE" in text
    assert "Sheet2" in text and called == [True]


def test_truncated_native_payload_is_not_accepted():
    with pytest.raises(ValueError):
        _native_payload(struct.pack("<I", 500) + b"\x02\0short")


@pytest.mark.parametrize("oversized", [False, True])
def test_rar_backend_passes_member_after_archive_and_limits_output(monkeypatch, oversized):
    import app.parsers.embedded as module
    monkeypatch.setattr(module.shutil, "which", lambda _: "/usr/bin/bsdtar")
    payload = workbook_bytes()
    class Process:
        def __init__(self, command, *, stdout, stderr):
            assert command[1] == "-xOf" and Path(command[2]).read_bytes() == b"rar-bytes"
            assert command[3:] == ["--", "评分/样例.xlsx"]
            self.stdout = BytesIO(payload + (b"oversized" if oversized else b""))
            self.returncode = None
        def __enter__(self):
            return self
        def __exit__(self, *args):
            self.stdout.close()
        def kill(self):
            self.returncode = -9
        def wait(self):
            self.returncode = self.returncode or 0
            return self.returncode
        def poll(self):
            return self.returncode
    monkeypatch.setattr(module.subprocess, "Popen", Process)
    info = SimpleNamespace(filename="评分/样例.xlsx", file_size=len(payload))
    if oversized:
        with pytest.raises(ValueError, match="RAR 解压失败或内容长度不符"):
            EmbeddedParser._rar_read(None, info, b"rar-bytes")
    else:
        assert EmbeddedParser._rar_read(None, info, b"rar-bytes") == payload


def test_rar_without_backend_explains_failure(monkeypatch):
    import app.parsers.embedded as module
    monkeypatch.setattr(module.shutil, "which", lambda _: None)
    def unavailable(_):
        raise module.rarfile.RarCannotExec("backend missing")
    with pytest.raises(ValueError, match="RAR 解压器不可用"):
        EmbeddedParser._rar_read(SimpleNamespace(read=unavailable), object(), b"")
