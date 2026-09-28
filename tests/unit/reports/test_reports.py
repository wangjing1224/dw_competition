"""JSON 原子写入、Word 生成和中文展示测试。"""

from docx import Document

from app.audit.sources import SourceLedger
from app.graph.state_updates import file_result_from_report
from app.models.documents import DocumentRef
from app.models.evidence import EvidenceLibrary
from app.models.findings import Finding
from app.models.results import AuditResult, FileAuditResult, ProjectMatch
from app.models.rules import RuleReportItem, present_rule_for_model
from app.reports.json_report import read_json_result, write_json_result
from app.reports.markdown_report import parse_report
from app.reports.word_report import generate_word_reports
from app.storage.file_store import FileStore
from tests.helpers.audit import audit_context, filled_report


def bid(file_id, applicable):
    document = DocumentRef(
        file_id=file_id,
        file_name=f"{file_id}.docx",
        source_path=f"{file_id}.docx",
        data_role="bid_document",
        parse_status="parsed",
        project_id="project-1",
        bidder_id="bidder-1",
        task_id="audit-1",
    )
    return FileAuditResult(
        file=document,
        is_applicable=applicable,
        project_match=ProjectMatch(
            status="matched" if applicable else "not_matched",
            reason="测试",
        ),
        exclusion_reason=None if applicable else "不匹配",
        overall_conclusion="manual_review" if applicable else "not_applicable",
    )


def test_word_is_generated_for_every_file_result(monkeypatch, temp_settings):
    monkeypatch.setattr("app.storage.file_store.get_settings", lambda: temp_settings)
    store = FileStore()
    result = AuditResult(
        task_id="audit-1",
        project_id="project-1",
        bidder_id="bidder-1",
        file_results=[bid("applicable", True), bid("ignored", False)],
        evidence_library=EvidenceLibrary(bidder_id="bidder-1"),
    )

    json_path = write_json_result("audit-1", result, store)
    word_paths = generate_word_reports(json_path, store)

    assert read_json_result(json_path, store) == result
    assert len(word_paths) == 2
    assert all(store.resolve_relative(path).exists() for path in word_paths)
    titles = [
        next(
            p.text
            for p in Document(store.resolve_relative(path)).paragraphs
            if p.style.name == "Title"
        )
        for path in word_paths
    ]
    assert titles == ["applicable审核报告", "ignored项目适用性说明"]


def test_model_title_survives_markdown_json_and_word(monkeypatch, temp_settings):
    monkeypatch.setattr("app.storage.file_store.get_settings", lambda: temp_settings)
    store = FileStore()
    context = audit_context()
    ledger = SourceLedger(context.parsed_documents)
    title = "红外热成像仪采购项目投标文件审核报告"
    decision, _ = parse_report(
        filled_report(context.rule_set, title=title),
        context.rule_set,
        ledger,
        final=True,
    )
    file_result = file_result_from_report(
        context.bid_files[0], context.bid_summaries["bid-1"], decision, ledger
    )
    result = AuditResult(
        task_id=context.task_id,
        project_id=context.project_id,
        bidder_id=context.bidder_id,
        rules=[
            RuleReportItem.model_validate(present_rule_for_model(r).model_dump())
            for r in context.rule_set.rules
        ],
        evidence_library=context.evidence_library,
        file_results=[file_result],
    )
    json_path = write_json_result(context.task_id, result, store)
    assert read_json_result(json_path, store).file_results[0].report_title == title
    word = Document(store.resolve_relative(generate_word_reports(json_path, store)[0]))
    assert [p.text for p in word.paragraphs if p.style.name == "Title"] == [title]
    assert "电力智能招标审核报告" not in [p.text for p in word.paragraphs]


def test_word_renders_markdown_and_enums_as_chinese(monkeypatch, temp_settings):
    monkeypatch.setattr("app.storage.file_store.get_settings", lambda: temp_settings)
    store = FileStore()
    file_result = bid("applicable", True).model_copy(
        update={
            "audit_summary": (
                "## 文件基本信息\n- **投标方**：某某科技有限公司\n- 已提交技术响应文件"
            ),
            "overall_conclusion": "non_compliant",
            "category_summaries": {"commercial": "商务响应存在缺失。"},
            "findings": [
                Finding(
                    rule_id="rule-1",
                    result="non_compliant",
                    summary="投标报价为空。",
                )
            ],
        }
    )
    result = AuditResult(
        task_id="audit-1",
        project_id="project-1",
        bidder_id="bidder-1",
        rules=[
            RuleReportItem(
                rule_id="rule-1",
                category="commercial",
                requirement="投标函应填写明确报价",
                priority="critical",
                verification="bid_only",
                reason="报价要求",
                source_file_ids=["rule-file"],
            )
        ],
        file_results=[file_result],
        evidence_library=EvidenceLibrary(bidder_id="bidder-1"),
    )
    json_path = write_json_result("audit-1", result, store)
    word_path = store.resolve_relative(generate_word_reports(json_path, store)[0])
    document = Document(word_path)
    text = "\n".join(
        [
            *(paragraph.text for paragraph in document.paragraphs),
            *(
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            ),
        ]
    )
    assert "文件基本信息" in text
    assert "投标方：某某科技有限公司" in text
    assert "商务响应" in text
    assert "不符合" in text
    assert "##" not in text
    assert "**" not in text
    assert "non_compliant" not in text
    assert "commercial" not in text
    assert "critical" not in text
