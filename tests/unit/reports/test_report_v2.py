"""新版契约、登记摘录与原生 Word 投影；所有文件仅在临时目录。"""

from datetime import UTC, datetime
from types import SimpleNamespace

import pytest
from docx import Document
from docx.oxml.ns import qn

from app.audit.sources import SourceLedger
from app.graph.state_updates import file_result_from_report
from app.models.results import AuditResult, ReportContext
from app.models.rules import RuleReportItem, present_rule_for_model
from app.reports.markdown_report import ReportDocument, parse_report, skeleton
from app.reports.presentation import report_markdown
from app.reports.word_report import _render_report
from app.tools.report_tools import ReportEdit
from tests.helpers.audit import audit_context, filled_report


def sample(count=2):
    context = audit_context()
    context.rule_set.rules = [
        context.rule_set.rules[0].model_copy(
            update={
                "rule_id": f"rule-{147 + i}",
                "requirement": f"第{i + 1}项技术要求" + "条件与单位不得遗漏。" * 60,
            }
        )
        for i in range(count)
    ]
    for parsed in context.parsed_documents.values():
        parsed.document.file_name = {
            "rule_source": "招标依据.md",
            "bid_document": "投标响应.md",
            "evidence_material": "证明材料.md",
        }[parsed.document.data_role]
    context.rule_set.source_files[0].file_name = "招标依据.md"
    ledger = SourceLedger(context.parsed_documents)
    text = filled_report(context.rule_set, version=2, conclusion="风险")
    text = text.replace(
        "无重点问题。",
        "### 问题 1：技术指标需复核\n\n关联规则："
        + "、".join(r.rule_id for r in reversed(context.rule_set.rules))
        + "\n\n问题概述：请结合原件复核参数。",
    )
    text = text.replace(
        "审核说明：分辨率满足要求",
        "审核说明：**需复核**实测记录。\n复核建议：核对原始测试记录。",
    )
    return context, ledger, text


def assemble(context, ledger, text):
    decision, missing = parse_report(
        text, context.rule_set, ledger, final=True, version=2
    )
    assert not missing
    f = file_result_from_report(
        context.bid_files[0],
        context.bid_summaries["bid-1"],
        decision,
        ledger,
        version=2,
        report_context=ReportContext(
            task_name="设备审核",
            rule_library_name="采购要求",
            evidence_library_name="供应商资料",
            finalized_at=datetime(2026, 10, 1, 4, 30, tzinfo=UTC),
        ),
    )
    return AuditResult(
        task_id=context.task_id,
        project_id=context.project_id,
        bidder_id=context.bidder_id,
        evidence_library=context.evidence_library,
        rule_source_files=[s.model_dump() for s in context.rule_set.source_files],
        rules=[
            RuleReportItem.model_validate(present_rule_for_model(r).model_dump())
            for r in context.rule_set.rules
        ],
        file_results=[f],
    )


def test_v2_complete_projection_quotes_and_native_word(tmp_path):
    context, ledger, text = sample()
    quote = "合同号 audit-2026，试验参数不得低于640×480。\n\n| 参数 | 要求 |\n|---|---|\n| 分辨率 | 640×480 |"
    ledger.chunk_read(
        SimpleNamespace(
            source=SimpleNamespace(file_id="tender-1", section="技术要求"),
            parse_version=context.parsed_documents["tender-1"].markdown_path,
            page_numbers=[3],
            content=quote,
            chunk_id="chunk-real",
        )
    )
    # 登记但未引用的片段不能自动附入报告。
    ledger.chunk_read(
        SimpleNamespace(
            source=SimpleNamespace(file_id="tender-1", section=None),
            parse_version=context.parsed_documents["tender-1"].markdown_path,
            page_numbers=[],
            content="未引用的正文",
            chunk_id="chunk-unused",
        )
    )
    result = assemble(context, ledger, text + "\n招标片段：chunk-real\n")
    f = result.file_results[0]
    assert f.category_summaries == {}
    assert f.suggestions == ["核对原始测试记录。"]
    assert set(f.sources) == {"chunk-real"}
    md = report_markdown(result, f)
    assert "2026年10月01日 12:30" in md
    assert "第4.1节、第4.2节" in md
    assert "[第4.1节](#41-技术响应风险)" in md
    assert md.count("核对原始测试记录。") == 3  # 重点中一次，两项明细各一次。
    assert "未引用的正文" not in md
    assert "\n> ".join(quote.splitlines()) in md
    assert all(
        key not in md
        for key in [
            "rule-147",
            "rule-148",
            "chunk-real",
            "project-1",
            "bidder-1",
            "audit-1",
        ]
    )
    for line in md.splitlines():
        if line.startswith("### 4."):
            assert len(line) < 40
    path = tmp_path / "report.docx"
    document = _render_report(result, f)
    document.save(path)
    document = Document(path)
    paragraphs = "\n".join(p.text for p in document.paragraphs)
    assert "**需复核**" not in paragraphs and "需复核" in paragraphs
    assert "audit-2026" in paragraphs  # 真实引文中的同形业务编号不删。
    assert all(
        key not in document.element.xml
        for key in ["rule-147", "rule-148", "chunk-real"]
    )
    assert len(document.tables) == 3  # 概况、统计、引文表格。
    bookmarks = document.element.findall(".//" + qn("w:bookmarkStart"))
    links = document.element.findall(".//" + qn("w:hyperlink"))
    assert {p.get(qn("w:name")) for p in bookmarks} == {
        "report-rule-4-1",
        "report-rule-4-2",
    }
    assert {p.get(qn("w:anchor")) for p in links} == {
        "report-rule-4-1",
        "report-rule-4-2",
    }
    assert all(
        t.rows[0]._tr.trPr.find(qn("w:tblHeader")) is not None for t in document.tables
    )


@pytest.mark.parametrize(
    "change",
    [
        lambda t: t.replace("实际发现：技术响应表载明参数。", ""),
        lambda t: t.replace("实际发现：", "实际发现：重复\n实际发现：", 1),
        lambda t: t.replace("关联规则：rule-147", "关联规则：rule-999"),
        lambda t: t.replace("关联规则：rule-147", "关联规则：rule-147、rule-147"),
        lambda t: t.replace("结论：风险", "结论：符合"),
        lambda t: t.replace("实际发现：技术响应表载明参数。", "实际发现：待填写。"),
        lambda t: t.replace("问题概述：", "问题概述：重复\n问题概述："),
        lambda t: t + "\n投标片段：chunk-invented\n",
        lambda t: t.replace("问题 1：技术指标需复核", "问题 1："),
        lambda t: t.replace("### rule-147", "### rule-999"),
    ],
)
def test_invalid_v2_draft_rejected(change):
    context, ledger, text = sample(1)
    with pytest.raises(ValueError):
        parse_report(change(text), context.rule_set, ledger, version=2)


@pytest.mark.asyncio
async def test_v2_invalid_draft_is_readable_and_editable_without_second_version(
    tmp_path,
):
    context, ledger, text = sample(1)
    document = ReportDocument(
        tmp_path / "draft.md", context.rule_set, ledger, version=2
    )
    decision, missing = parse_report(
        skeleton(context.rule_set, 2), context.rule_set, ledger, version=2
    )
    assert decision is None and "重点问题" in missing
    with pytest.raises(ValueError):
        document.finish()
    await document.apply("write", text, None, "save")
    invalid = text + "\n资料片段：fake\n"
    receipt = await document.apply("write", invalid, None, "invalid")
    assert "草稿已保存" in receipt and "校验未通过" in receipt
    assert "rule-147的资料片段" in receipt
    assert document.read() == invalid
    assert "资料片段：fake" in await document.apply("read", None, None, "read")
    assert document.missing()  # 构造下一轮请求不能因内容校验中断。
    with pytest.raises(ValueError, match="未实际读取"):
        document.finish()
    await document.apply(
        "edit", None, [ReportEdit(old_text="\n资料片段：fake\n", new_text="")], "repair"
    )
    assert document.finish().highlights[0].rule_ids == ["rule-147"]
    assert list(tmp_path.iterdir()) == [document.path]


def test_56_rules_complete_no_forced_highlights_or_truncation():
    context, ledger, _ = sample(56)
    text = filled_report(
        context.rule_set, version=2, reason="合同号 audit-2026 的参数符合要求。"
    )
    result = assemble(context, ledger, text)
    md = report_markdown(result, result.file_results[0])
    assert len(result.file_results[0].findings) == 56
    assert "无重点问题。" in md and "### 4.56 技术响应｜符合" in md
    assert context.rule_set.rules[-1].requirement in md
    assert "| 56项 | 56项 | 0项 | 0项 | 0项 |" in md


@pytest.mark.asyncio
async def test_missing_fixed_category_is_restored_before_saving(tmp_path):
    context, ledger, _ = sample(3)
    context.rule_set.rules[0].category = "format"
    text = filled_report(context.rule_set, version=2)
    incomplete = text.replace("## 技术响应（technical）\n\n", "")
    report = ReportDocument(tmp_path / "draft.md", context.rule_set, ledger, version=2)
    receipt = await report.apply("write", incomplete, None, "save")
    assert "已按冻结规则补齐" in receipt
    assert report.read() == text
    assert len(report.finish().categories) == 2
    # 显式重导出同样可恢复确定的固定标题，不依赖模型或数据库。
    decision, _ = parse_report(
        incomplete, context.rule_set, ledger, final=True, version=2
    )
    assert len(decision.categories) == 2


@pytest.mark.parametrize(
    "change",
    [
        lambda t: t.replace("### rule-148", "### rule-999"),
        lambda t: t.replace("### rule-148", "### rule-147"),
        lambda t: (
            t.replace("### rule-147", "### temp")
            .replace("### rule-148", "### rule-147")
            .replace("### temp", "### rule-148")
        ),
        lambda t: t[: t.index("### rule-148")],
    ],
)
def test_category_repair_does_not_guess_unknown_missing_repeated_or_reordered_rules(
    change,
):
    from app.reports.report_parser_v2 import restore_category_headings

    context, ledger, _ = sample(2)
    text = change(
        filled_report(context.rule_set, version=2).replace(
            "## 技术响应（technical）\n\n", ""
        )
    )
    assert restore_category_headings(text, context.rule_set) == (text, [])
    with pytest.raises(ValueError, match="类别或规则集合"):
        parse_report(text, context.rule_set, ledger, version=2, final=True)


def test_known_ids_convert_only_prose_using_frozen_names_and_rule_order():
    context, ledger, text = sample(2)
    names = ReportContext(
        task_name="十月审核",
        rule_library_name="设备要求",
        evidence_library_name="证明材料库",
    )
    text = text.replace(
        "技术响应表载明参数。",
        "bid-1载明参数；参见rule-148，按project-1核对bidder-1，任务audit-1；合同号audit-2026、rule-1480。",
    )
    decision, _ = parse_report(
        text, context.rule_set, ledger, version=2, final=True, report_context=names
    )
    observation = decision.categories[0].rules[0].summary.observation
    assert (
        observation
        == "投标响应.md载明参数；参见第4.2节，按设备要求核对证明材料库，任务十月审核；合同号audit-2026、rule-1480。"
    )
    assert decision.categories[0].rules[0].rule_id == "rule-147"
    assert decision.highlights[0].rule_ids == ["rule-148", "rule-147"]
    assert "bid-1" in text  # 内部原稿不被名称转换修改。


def test_known_source_references_convert_without_rewriting_actual_quotes_or_file_names():
    context, ledger, text = sample(1)
    quote = "合同号为rule-147，实际设备编号为bid-1。"
    ledger.chunk_read(
        SimpleNamespace(
            source=SimpleNamespace(file_id="tender-1", section="技术要求"),
            parse_version=context.parsed_documents["tender-1"].markdown_path,
            page_numbers=[3],
            content=quote,
            chunk_id="chunk-real",
        )
    )
    ledger.documents["bid-1"].document.file_name = "bid-1投标.md"
    text = text.replace(
        "技术响应表载明参数。",
        f"摘自chunk-real：{quote}参见bid-1投标.md。另见rule-147。",
    )
    decision, _ = parse_report(text, context.rule_set, ledger, final=True, version=2)
    observation = decision.categories[0].rules[0].summary.observation
    assert (
        observation
        == f"摘自招标依据.md · 技术要求 · 页 3：{quote}参见bid-1投标.md。另见第4.1节。"
    )
    assert ledger.sources["chunk-real"].quote == quote


def test_ambiguous_known_file_reference_is_not_guessed():
    context, ledger, text = sample(1)
    context.rule_set.source_files[0].file_name = "另一名称.md"
    with pytest.raises(ValueError, match="存在歧义"):
        parse_report(
            text.replace("技术响应表载明参数。", "参见tender-1。"),
            context.rule_set,
            ledger,
            version=2,
        )


def test_markdown_native_lists_links_and_no_remote_images():
    from app.reports.word_markdown import render_markdown

    document = Document()
    render_markdown(
        document,
        "1. **第一项**\n2. 第二项\n\n- 一条建议\n\n[参考](https://example.org)\n\n![仅保留图名](https://example.org/private.png)\n\n<script>alert(1)</script>",
    )
    assert (
        len(
            [
                p
                for p in document.paragraphs
                if p._p.pPr is not None and p._p.pPr.numPr is not None
            ]
        )
        == 3
    )
    assert any(
        r.bold and r.text == "第一项" for p in document.paragraphs for r in p.runs
    )
    assert "<script>alert(1)</script>" in "\n".join(p.text for p in document.paragraphs)
    assert not document.inline_shapes
    assert any(
        r.target_ref == "https://example.org" for r in document.part.rels.values()
    )


def test_image_reference_without_quote_does_not_invent_ocr_or_page():
    from app.models.documents import ParsedImage

    context, ledger, text = sample(1)
    path = "bid-1/images/scan.png"
    context.parsed_documents["bid-1"].sections[0].blocks[0].images = [
        ParsedImage(path=path, mime_type="image/png", width=100, height=100)
    ]
    ledger.image_read(path)
    result = assemble(context, ledger, text + f"\n图片来源：{path}\n")
    assert not result.file_results[0].sources[path].quote
    md = report_markdown(result, result.file_results[0])
    assert "**投标图片依据**" in md and "投标响应.md" in md
    assert "解析原文摘录" not in md and "页 " not in md
    assert path not in md


def test_wrong_source_role_and_version_rejected():
    context, ledger, text = sample(1)
    hit = SimpleNamespace(
        source=SimpleNamespace(file_id="tender-1", section="技术要求"),
        parse_version="wrong-version",
        page_numbers=[],
        content="参数100",
        chunk_id="chunk-read",
    )
    with pytest.raises(ValueError, match="版本"):
        ledger.chunk_read(hit)
    hit.parse_version = context.parsed_documents["tender-1"].markdown_path
    ledger.chunk_read(hit)
    with pytest.raises(ValueError, match="类型不符"):
        assemble(context, ledger, text + "\n投标片段：chunk-read\n")


def test_mixed_results_counts_and_conclusion_priority():
    context, ledger, text = sample(4)
    for index, conclusion in enumerate(["不符合", "符合", "人工复核", "风险"]):
        text = text.replace(
            f"### rule-{147 + index}\n\n结论：风险",
            f"### rule-{147 + index}\n\n结论：{conclusion}",
        )
    result = assemble(context, ledger, text)
    f = result.file_results[0]
    assert f.overall_conclusion == "non_compliant"
    md = report_markdown(result, f)
    assert "| 4项 | 1项 | 1项 | 1项 | 1项 |" in md
    assert "**涉及结论：** 不符合、符合、人工复核、风险" in md
