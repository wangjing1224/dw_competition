"""根据已校验 JSON 生成正式中文 Word 审核报告。"""

from __future__ import annotations

import os
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor

from app.models.results import AuditResult, FileAuditResult
from app.models.rules import RuleReportItem
from app.reports.json_report import read_json_result
from app.reports.styles import (
    ACCENT,
    DARK,
    MUTED,
    apply_document_styles,
    category_label,
    conclusion_label,
    finding_label,
    match_label,
    priority_label,
)
from app.storage.file_store import FileStore

CATEGORY_ORDER = {
    "format": 0,
    "qualification": 1,
    "commercial": 2,
    "technical": 3,
    "performance_capacity": 4,
    "consistency_risk": 5,
}
PRIORITY_ORDER = {"critical": 0, "important": 1, "normal": 2}
FINDING_ORDER = {
    "non_compliant": 0,
    "risk": 1,
    "manual_review": 2,
    "compliant": 3,
}


def _safe_stem(file_name: str) -> str:
    stem = Path(file_name).stem
    return re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", stem)[:80] or "投标文件"


def _ordered_findings(
    file_result: FileAuditResult,
    rules: dict[str, RuleReportItem],
):
    return sorted(
        file_result.findings,
        key=lambda finding: (
            CATEGORY_ORDER.get(str(rules[finding.rule_id].category), 99),
            PRIORITY_ORDER.get(str(rules[finding.rule_id].priority), 99),
            FINDING_ORDER.get(str(finding.result), 99),
            finding.rule_id,
        ),
    )


def _set_run(
    run,
    *,
    chinese: str = "宋体",
    latin: str = "Arial",
    size: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
) -> None:
    from docx.oxml.ns import qn

    run.font.name = latin
    run._element.rPr.rFonts.set(qn("w:ascii"), latin)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), latin)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), chinese)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def _add_label_value(document, label: str, value: str, *, style=None) -> None:
    paragraph = document.add_paragraph(style=style)
    label_run = paragraph.add_run(f"{label}：")
    _set_run(label_run, chinese="黑体", bold=True, color=DARK)
    value_run = paragraph.add_run(value or "未记录")
    _set_run(value_run)


def _add_title_block(
    document,
    result: AuditResult,
    file_result: FileAuditResult,
) -> None:
    kicker = document.add_paragraph()
    kicker.paragraph_format.space_before = Pt(18)
    kicker.paragraph_format.space_after = Pt(6)
    run = kicker.add_run("自动审核结果")
    _set_run(run, chinese="黑体", size=10, bold=True, color=ACCENT)

    title = document.add_paragraph(style="Title")
    title.add_run(file_result.display_title)
    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    run = subtitle.add_run(file_result.file.file_name)
    _set_run(run, chinese="黑体", size=14, color=MUTED)

    metadata = (
        ("审核任务编号", result.task_id),
        ("项目编号", result.project_id),
        ("投标方编号", result.bidder_id),
        ("生成时间", datetime.now().strftime("%Y年%m月%d日 %H:%M")),
    )
    for label, value in metadata:
        _add_label_value(document, label, value, style="Report Metadata")

    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(6)


def _add_summary_markdown(document, text: str) -> None:
    """渲染摘要中实际出现的最小 Markdown 子集。"""

    paragraph_buffer: list[str] = []

    def flush() -> None:
        if paragraph_buffer:
            value = re.sub(r"\*\*(.+?)\*\*", r"\1", " ".join(paragraph_buffer))
            document.add_paragraph(value.strip())
            paragraph_buffer.clear()

    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            flush()
            continue
        heading = re.match(r"^(#{2,3})\s+(.+)$", line)
        if heading:
            flush()
            level = 2 if len(heading.group(1)) == 2 else 3
            document.add_heading(
                re.sub(r"\*\*(.+?)\*\*", r"\1", heading.group(2)),
                level=level,
            )
            continue
        field = re.match(r"^-\s+\*\*(.+?)\*\*[：:]\s*(.*)$", line)
        if field:
            flush()
            _add_label_value(document, field.group(1), field.group(2))
            continue
        bullet = re.match(r"^-\s+(.+)$", line)
        if bullet:
            flush()
            value = re.sub(r"\*\*(.+?)\*\*", r"\1", bullet.group(1))
            document.add_paragraph(value, style="List Bullet")
            continue
        paragraph_buffer.append(line)
    flush()


def _finding_sources(finding, file_result) -> list[str]:
    values: list[str] = []
    for label, chunk_ids in (
        ("招标原文", finding.tender_chunk_ids),
        ("投标原文", finding.bid_chunk_ids),
        ("支撑资料", finding.evidence_chunk_ids),
        ("图片", finding.image_paths),
    ):
        for chunk_id in chunk_ids:
            source = file_result.sources.get(chunk_id)
            if source is None:
                raise ValueError(f"报告缺少引用快照：{chunk_id}")
            location = " / ".join(
                value
                for value in (
                    source.file_name,
                    source.section,
                    "页码：" + "、".join(map(str, source.page_numbers))
                    if source.page_numbers
                    else "",
                )
                if value
            )
            values.append(
                f"{label}：{location or source.source_path}；{chunk_id}"
                + (f"\n摘录：{source.quote}" if source.quote else "")
            )
    return values


def _add_finding_details(
    document,
    result: AuditResult,
    file_result: FileAuditResult,
    rules: dict[str, RuleReportItem],
) -> None:
    document.add_heading("三、分类审核", level=1)
    grouped: dict[str, list] = defaultdict(list)
    for finding in _ordered_findings(file_result, rules):
        grouped[str(rules[finding.rule_id].category)].append(finding)
    for category in sorted(
        grouped,
        key=lambda value: CATEGORY_ORDER.get(value, 99),
    ):
        document.add_heading(category_label(category), level=2)
        _add_summary_markdown(
            document, file_result.category_summaries.get(category, "")
        )
        for finding in grouped[category]:
            rule = rules[finding.rule_id]
            heading = document.add_heading(level=3)
            heading.add_run(
                f"{finding.rule_id}｜{priority_label(rule.priority)}｜"
                f"{finding_label(finding.result)}"
            )
            _add_label_value(document, "审核要求", rule.requirement)
            _add_label_value(document, "审核结论", finding_label(finding.result))
            document.add_paragraph("审核说明", style="Finding Label")
            _add_summary_markdown(document, finding.summary)
            sources = _finding_sources(finding, file_result)
            if sources:
                label = document.add_paragraph(style="Finding Label")
                label.add_run("引用依据")
                for value in sources:
                    document.add_paragraph(value, style="List Bullet")


def _add_suggestions(
    document,
    file_result: FileAuditResult,
) -> None:
    document.add_heading("四、整改与复核建议", level=1)
    if not file_result.suggestions:
        document.add_paragraph("暂无额外整改建议。")
        return
    for suggestion in file_result.suggestions:
        document.add_paragraph(suggestion, style="List Bullet")


def _render_report(
    result: AuditResult,
    file_result: FileAuditResult,
):
    rules = {rule.rule_id: rule for rule in result.rules}
    document = Document()
    apply_document_styles(document)
    _add_title_block(document, result, file_result)
    document.add_heading("一、项目匹配情况", level=1)
    _add_label_value(
        document, "匹配状态", match_label(file_result.project_match.status)
    )
    _add_summary_markdown(document, file_result.project_match.reason)

    if not file_result.is_applicable:
        document.add_heading("未进入审核", level=1)
        document.add_paragraph(file_result.exclusion_reason or "未记录具体原因")
        for suggestion in file_result.suggestions:
            document.add_paragraph(suggestion, style="List Bullet")
        return document

    document.add_heading("二、重点结论与待核实事项", level=1)
    _add_label_value(
        document, "总体结论", conclusion_label(file_result.overall_conclusion)
    )
    _add_summary_markdown(document, file_result.audit_summary or "")
    _add_finding_details(document, result, file_result, rules)
    if file_result.suggestions:
        _add_suggestions(document, file_result)
    return document


def generate_word_reports(
    json_path: str,
    file_store: FileStore | None = None,
) -> list[str]:
    store = file_store or FileStore()
    result = read_json_result(json_path, store)
    directory = store.task_output_dir(result.task_id)
    paths: list[str] = []
    for file_result in result.file_results:
        path = directory / (
            f"{_safe_stem(file_result.file.file_name)}_{file_result.file.file_id[:12]}"
            "_审核报告.docx"
        )
        temporary = path.with_suffix(".docx.tmp")
        document = _render_report(result, file_result)
        document.save(temporary)
        os.replace(temporary, path)
        paths.append(store.relative_to_data_root(path))
    return paths
