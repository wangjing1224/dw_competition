"""Word 审核报告的中文展示映射和统一版式。

设计基线采用 ``standard_business_brief``，并使用命名覆盖
``a4_chinese_audit_report`` 将页面调整为更适合国内正式报告的 A4。
"""

from docx.document import Document as WordDocument
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

CATEGORY_LABELS = {
    "format": "格式与文件完整性",
    "qualification": "资格条件",
    "commercial": "商务响应",
    "technical": "技术响应",
    "performance_capacity": "履约与业绩能力",
    "consistency_risk": "一致性与风险",
}
PRIORITY_LABELS = {
    "critical": "关键",
    "important": "重要",
    "normal": "一般",
}
FINDING_LABELS = {
    "compliant": "符合",
    "non_compliant": "不符合",
    "manual_review": "人工复核",
}
MATCH_LABELS = {
    "matched": "匹配",
    "not_matched": "不匹配",
    "uncertain": "无法确定",
}
CONCLUSION_LABELS = {
    "compliant": "总体符合",
    "non_compliant": "存在不符合项",
    "manual_review": "需人工复核",
    "not_applicable": "不适用",
}
EVIDENCE_TYPE_LABELS = {
    "business_license": "营业执照",
    "certificate": "资质与认证证书",
    "contract": "合同与业绩材料",
    "invoice": "发票与交易凭证",
    "test_report": "检测与试验报告",
    "financial_report": "财务报告",
    "authorization": "授权材料",
    "other": "其他资料",
}

ACCENT = "1F4E78"
DARK = "17212B"
MUTED = "667085"
LIGHT_FILL = "F2F4F7"
ACCENT_FILL = "EAF2F8"
RISK_FILL = "FDECEC"
MANUAL_FILL = "F3F0FF"
POSITIVE_FILL = "EAF6EE"
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def category_label(value: object) -> str:
    return CATEGORY_LABELS.get(str(value), str(value))


def priority_label(value: object) -> str:
    return PRIORITY_LABELS.get(str(value), str(value))


def finding_label(value: object) -> str:
    return FINDING_LABELS.get(str(value), str(value))


def match_label(value: object) -> str:
    return MATCH_LABELS.get(str(value), str(value))


def conclusion_label(value: object) -> str:
    return CONCLUSION_LABELS.get(str(value), str(value))


def evidence_type_label(value: object) -> str:
    return EVIDENCE_TYPE_LABELS.get(str(value), str(value))


def _set_style_font(style, chinese: str, latin: str, size: float) -> None:
    style.font.name = latin
    style.font.size = Pt(size)
    fonts = style._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), latin)
    fonts.set(qn("w:hAnsi"), latin)
    fonts.set(qn("w:eastAsia"), chinese)


def _set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, *, top=100, start=120, bottom=100, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    """显式设置 tblW、tblInd、tblGrid 和每个 tcW。"""

    if sum(widths) != TABLE_WIDTH_DXA:
        raise ValueError("Word 表格列宽之和必须等于 9360 DXA")
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def format_table_header(row) -> None:
    _set_repeat_table_header(row)
    for cell in row.cells:
        shade_cell(cell, LIGHT_FILL)
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.bold = True
                run.font.color.rgb = RGBColor.from_string(DARK)


def result_fill(value: object) -> str:
    return {
        "compliant": POSITIVE_FILL,
        "non_compliant": RISK_FILL,
        "manual_review": MANUAL_FILL,
    }.get(str(value), LIGHT_FILL)


def apply_document_styles(document: WordDocument) -> None:
    """应用正式中文审核报告样式和页眉页脚。"""

    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)
    section.header_distance = Cm(1)
    section.footer_distance = Cm(1)

    normal = document.styles["Normal"]
    _set_style_font(normal, "宋体", "Arial", 10.5)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.35

    title = document.styles["Title"]
    _set_style_font(title, "黑体", "Arial", 21)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(DARK)
    title.paragraph_format.space_before = Pt(6)
    title.paragraph_format.space_after = Pt(8)
    title.paragraph_format.keep_with_next = True

    heading_tokens = {
        "Heading 1": (16, 16, 8),
        "Heading 2": (13, 12, 6),
        "Heading 3": (11.5, 8, 4),
    }
    for style_name, (size, before, after) in heading_tokens.items():
        style = document.styles[style_name]
        _set_style_font(style, "黑体", "Arial", size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(ACCENT)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    list_style = document.styles["List Bullet"]
    _set_style_font(list_style, "宋体", "Arial", 10.5)
    list_style.paragraph_format.left_indent = Cm(0.74)
    list_style.paragraph_format.first_line_indent = Cm(-0.37)
    list_style.paragraph_format.space_after = Pt(4)
    list_style.paragraph_format.line_spacing = 1.2

    if "Report Metadata" not in document.styles:
        metadata = document.styles.add_style(
            "Report Metadata",
            WD_STYLE_TYPE.PARAGRAPH,
        )
    else:
        metadata = document.styles["Report Metadata"]
    _set_style_font(metadata, "宋体", "Arial", 9.5)
    metadata.font.color.rgb = RGBColor.from_string(MUTED)
    metadata.paragraph_format.space_after = Pt(2)
    metadata.paragraph_format.line_spacing = 1.1

    if "Finding Label" not in document.styles:
        label = document.styles.add_style(
            "Finding Label",
            WD_STYLE_TYPE.PARAGRAPH,
        )
    else:
        label = document.styles["Finding Label"]
    _set_style_font(label, "黑体", "Arial", 10.5)
    label.font.bold = True
    label.font.color.rgb = RGBColor.from_string(DARK)
    label.paragraph_format.space_before = Pt(4)
    label.paragraph_format.space_after = Pt(2)
    label.paragraph_format.keep_with_next = True

    header = section.header
    header_paragraph = header.paragraphs[0]
    header_paragraph.text = "招投标智能审核 · 辅助审核报告"
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header_paragraph.runs:
        run.font.name = "Arial"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string(MUTED)

    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer_paragraph.add_run("第 ")
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.append(field_begin)
    run._r.append(instruction)
    run._r.append(field_end)
    footer_paragraph.add_run(" 页")
    for footer_run in footer_paragraph.runs:
        footer_run.font.name = "Arial"
        footer_run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        footer_run.font.size = Pt(9)
        footer_run.font.color.rgb = RGBColor.from_string(MUTED)
