"""Markdown 和解析摘录表格到原生 Word；不执行 HTML 或请求外部图片。"""

import math
import re
from urllib.parse import unquote, urlsplit

from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree
from markdown_it import MarkdownIt
from markdown_it.tree import SyntaxTreeNode

from app.reports.presentation import heading_anchor
from app.reports.styles import (
    ACCENT,
    LIGHT_FILL,
    TABLE_WIDTH_DXA,
    format_table_header,
    set_table_geometry,
)


def _excerpt_table_block(state, start, end, silent):
    """仅识别完整的 HTML 表格块；其余 HTML 与代码仍按字面处理。"""
    pos = state.bMarks[start] + state.tShift[start]
    if state.is_code_block(start) or not re.match(
        r"<table(?:\s|>)", state.src[pos:state.eMarks[start]], re.I
    ):
        return False
    for line in range(start, end):
        if state.sCount[line] < state.blkIndent:
            break
        value = state.src[state.bMarks[line] + state.tShift[line]:state.eMarks[line]]
        if re.search(r"</table\s*>", value, re.I):
            if not silent:
                state.line = line + 1
                token = state.push("excerpt_table", "", 0)
                token.map = [start, state.line]
                token.content = state.getLines(start, state.line, state.blkIndent, True)
            return True
    return False


def _html_cell(paragraph, element, *, bold=False, italic=False):
    """只投影文字、换行和强调；忽略样式、脚本及图片地址。"""
    def add(text):
        if text:
            run = paragraph.add_run(text)
            run.bold, run.italic = bold, italic

    add(element.text)
    for child in element:
        if child.tag == "br":
            add("\n")
        elif child.tag == "img":
            add(child.get("alt", ""))
        elif isinstance(child.tag, str) and child.tag not in {"script", "style"}:
            if child.tag in {"p", "div", "tr", "td", "th"} and paragraph.text:
                add("\n")
            _html_cell(
                paragraph, child,
                bold=bold or child.tag in {"b", "strong"},
                italic=italic or child.tag in {"i", "em"},
            )
        add(child.tail)


def _render_excerpt_table(container, html, page_layout=None):
    # lxml 已由 python-docx 使用；仅在内存中解析，不加载 URL 或外部资源。
    root = etree.HTML(html, parser=etree.HTMLParser(no_network=True))
    rows = root.xpath(
        "./body/table/tr | ./body/table/thead/tr | "
        "./body/table/tbody/tr | ./body/table/tfoot/tr"
    )
    placements, occupied = [], set()
    try:
        for r, row in enumerate(rows):
            c = 0
            for cell in row.xpath("./td | ./th"):
                while (r, c) in occupied:
                    c += 1
                rs, cs = int(cell.get("rowspan", 1)), int(cell.get("colspan", 1))
                # Word 表格最多 63 列；异常跨度不扩张成巨大的内存网格。
                if not (1 <= rs <= len(rows) - r and 1 <= cs <= 63 - c):
                    raise ValueError("无效表格跨度")
                area = {(y, x) for y in range(r, r + rs) for x in range(c, c + cs)}
                if occupied & area:
                    raise ValueError("表格单元格重叠")
                occupied.update(area)
                placements.append((r, c, rs, cs, cell))
                c += cs
        columns = max(c for _, c in occupied) + 1
    except ValueError:
        # 非规范摘录仍保留可读内容，不伪造合并关系或令整份报告导出失败。
        container.add_paragraph("原文表格结构无法还原，以下按文字列出：")
        _html_cell(container.add_paragraph(), root)
        return

    # 只采信 th/thead；未知表格不再把空行或首条记录当作列名。
    header_rows = {
        i for i, row in enumerate(rows)
        if row.getparent().tag == "thead"
        or (row.xpath("./th") and not row.xpath("./td"))
    }
    first_header = min(header_rows, default=0)
    # 空数据单元格不删除。只清理首尾整行留白，以及表头前的排版留白；
    # 纵向合并占用的行必须保留，否则会改变单元格归属。
    nonempty = {
        r for r, _, _, _, cell in placements
        if "".join(cell.itertext()).strip() or cell.xpath(".//img")
    }
    protected = {y for r, _, rs, _, _ in placements if rs > 1 for y in range(r, r + rs)}
    if not nonempty and not protected:
        return
    first, last = min(nonempty | protected), max(nonempty | protected)
    kept = [
        r for r in range(len(rows))
        if r in protected or r in nonempty or (max(first, first_header) <= r <= last)
    ]
    row_map = {r: i for i, r in enumerate(kept)}
    placements = [(row_map[r], c, rs, cs, cell) for r, c, rs, cs, cell in placements if r in row_map]
    header_rows = {row_map[r] for r in header_rows if r in row_map}
    first_header = min(header_rows, default=0)
    total_width = page_layout(columns > 8) if page_layout else TABLE_WIDTH_DXA

    # 工作表标题/项目说明在列名之前，作为表前文字展示，不挤进第一列；
    # 遇到跨越分界的纵向合并则不拆，优先保证结构准确。
    if first_header and not any(r < first_header < r + rs for r, _, rs, _, _ in placements):
        for row in range(first_header):
            cells = [
                cell for r, _, _, _, cell in placements
                if r == row and ("".join(cell.itertext()).strip() or cell.xpath(".//img"))
            ]
            if cells:
                p = container.add_paragraph()
                for i, cell in enumerate(cells):
                    if i:
                        p.add_run(" · ")
                    _html_cell(p, cell)
                p.paragraph_format.keep_with_next = True
                p.paragraph_format.space_after = Pt(4)
        placements = [
            (r - first_header, c, rs, cs, cell)
            for r, c, rs, cs, cell in placements if r >= first_header
        ]
        kept = kept[first_header:]
        header_rows = {r - first_header for r in header_rows}

    table = container.add_table(rows=len(kept), cols=columns)
    table.style = "Table Grid"
    # 为序号、数量等短列少分宽度，给要求/响应正文留空间；总宽保持报告版心。
    weights = [4] * columns
    for _, c, _, cs, cell in placements:
        weight = min(40, max((len(t) for t in cell.itertext()), default=0) / cs)
        for index in range(c, c + cs):
            weights[index] = max(weights[index], weight)
    source_widths = root.xpath("./body/table/colgroup/col/@data-width")
    try:
        source_widths = [float(w) for w in source_widths]
        if len(source_widths) == columns and all(math.isfinite(w) and 0 < w <= 4096 for w in source_widths):
            weights = source_widths
    except ValueError:
        pass  # 不可信/旧格式元数据不影响文字和网格渲染。
    minimum = min(720, total_width // columns)
    spare = total_width - minimum * columns
    widths = [minimum + int(spare * w / sum(weights)) for w in weights]
    widths[-1] += total_width - sum(widths)
    set_table_geometry(table, widths, total_width=total_width)
    for r, c, rs, cs, source in placements:
        cell = table.cell(r, c)
        if rs > 1 or cs > 1:
            cell = cell.merge(table.cell(r + rs - 1, c + cs - 1))
        p = cell.paragraphs[0]
        _html_cell(p, source)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        for run in p.runs:
            run.font.size = Pt(9 if columns >= 6 else 10)
    # 合并表头跨行时，重复整个表头区域，避免跨页时拆断纵向合并。
    header_end = 0
    while header_end in header_rows:
        header_end += 1
    for r, _, rs, _, _ in placements:
        if r < header_end:
            header_end = max(header_end, r + rs)
    for row in table.rows[:header_end]:
        format_table_header(row)
    for row in table.rows:
        # 普通行尽量整行跨页，避免续页只剩某个长单元格的两行尾句。
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def _numbering(document, ordered, start=1):
    root = document.part.numbering_part.element
    abstract_id = (
        max(
            [
                int(x.get(qn("w:abstractNumId")))
                for x in root.findall(qn("w:abstractNum"))
            ]
            + [-1]
        )
        + 1
    )
    num_id = (
        max([int(x.get(qn("w:numId"))) for x in root.findall(qn("w:num"))] + [0]) + 1
    )
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    for tag, val in (
        ("start", str(start)),
        ("numFmt", "decimal" if ordered else "bullet"),
        ("lvlText", "%1." if ordered else "•"),
    ):
        elem = OxmlElement("w:" + tag)
        elem.set(qn("w:val"), val)
        level.append(elem)
    abstract.append(level)
    root.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), str(abstract_id))
    num.append(ref)
    root.append(num)
    return num_id


def _inline(paragraph, tokens, *, bold=False, italic=False, link=None, anchors=None):
    for t in tokens:
        if t.type == "strong_open":
            bold = True
        elif t.type == "strong_close":
            bold = False
        elif t.type == "em_open":
            italic = True
        elif t.type == "em_close":
            italic = False
        elif t.type == "link_open":
            link = t.attrGet("href")
        elif t.type == "link_close":
            link = None
        elif t.type in {
            "text",
            "code_inline",
            "html_inline",
            "image",
            "softbreak",
            "hardbreak",
        }:
            text = "\n" if t.type in {"softbreak", "hardbreak"} else t.content
            run = paragraph.add_run(text)
            run.bold, run.italic = bold, italic
            if t.type == "code_inline":
                run.font.name = "Courier New"
            if link and (
                link.startswith("#")
                or urlsplit(link).scheme in {"http", "https", "mailto"}
            ):
                hyperlink = OxmlElement("w:hyperlink")
                if link.startswith("#"):
                    anchor = unquote(link[1:])
                    hyperlink.set(qn("w:anchor"), (anchors or {}).get(anchor, anchor))
                else:
                    hyperlink.set(
                        qn("r:id"),
                        paragraph.part.relate_to(link, RT.HYPERLINK, is_external=True),
                    )
                run.font.color.rgb = RGBColor.from_string(ACCENT)
                hyperlink.append(run._r)
                paragraph._p.append(hyperlink)


def render_markdown(document, text):
    parser = MarkdownIt("commonmark", {"html": False}).enable("table")
    parser.block.ruler.before(
        "html_block", "excerpt_table", _excerpt_table_block,
        {"alt": ["paragraph", "reference", "blockquote"]},
    )
    root = SyntaxTreeNode(parser.parse(text))
    nodes = root.children
    caption_nodes = set()
    for i in range(len(nodes) - 2):
        label, source, following = nodes[i:i + 3]
        if (
            label.type == source.type == "paragraph"
            and re.fullmatch(r"\*\*[^*]+\*\*", label.children[0].token.content)
            and (following.type == "excerpt_table" or (
                following.type == "blockquote" and following.children
                and following.children[0].type == "excerpt_table"
            ))
        ):
            caption_nodes.update((id(label), id(source)))
    bookmark_count = len(document.element.findall(".//" + qn("w:bookmarkStart")))
    anchors = {}
    original = document.sections[-1]
    page_width, page_height = sorted((original.page_width, original.page_height))
    portrait_table_width = min(
        TABLE_WIDTH_DXA,
        int((page_width - original.left_margin - original.right_margin) / 635) - 120,
    )

    def page_layout(landscape=False):
        """宽摘录独占横向节；下段正文恢复竖向，不在末尾制造空白页。"""
        orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
        current = document.sections[-1]
        if current.orientation != orientation:
            # 把紧接表格的“摘录标题 + 来源”一起带到横向页，避免留下孤立标题。
            captions = []
            body = document._element.body
            if len(body) >= 3:
                a, b = body[-3], body[-2]
                if a.tag == b.tag == qn("w:p") and a.xpath("./w:pPr/w:keepNext"):
                    captions = [a, b]
            for p in captions:
                body.remove(p)
            if len(body) > 1:
                current = document.add_section(WD_SECTION_START.NEW_PAGE)
            current.orientation = orientation
            current.page_width = page_height if landscape else page_width
            current.page_height = page_width if landscape else page_height
            for p in captions:
                body.insert_element_before(p, "w:sectPr")
        return (
            int((current.page_width - current.left_margin - current.right_margin) / 635) - 120
            if landscape else portrait_table_width
        )
    for title in re.findall(r"^### (4\.\d+ .+)$", text, re.MULTILINE):
        anchors[heading_anchor(title)] = "report-rule-" + title.split(" ", 1)[
            0
        ].replace(".", "-")

    def paragraph(container, node, quote=False, numbering=None, depth=0):
        # 表前标题和来源的页面方向由随后的表格决定，连续宽表不插入空竖页。
        if id(node) not in caption_nodes:
            page_layout()
        p = container.add_paragraph()
        if id(node) in caption_nodes:
            p.paragraph_format.keep_with_next = True
        if node.type == "heading" and not quote:
            level = int(node.tag[1:])
            p.style = "Title" if level == 1 else f"Heading {min(level - 1, 3)}"
        if numbering:
            props = p._p.get_or_add_pPr().get_or_add_numPr()
            props.get_or_add_ilvl().val = 0
            props.get_or_add_numId().val = numbering
            p.paragraph_format.left_indent = Cm(0.65 * max(1, depth))
            p.paragraph_format.first_line_indent = Cm(-0.35)
        for child in node.children or []:
            if child.type == "inline":
                _inline(p, child.token.children or [], anchors=anchors)
        if node.type in {"fence", "code_block"}:
            p.add_run(node.content).font.name = "Courier New"
        if quote:
            p.paragraph_format.left_indent = Cm(0.4)
            p.paragraph_format.right_indent = Cm(0.2)
            p.paragraph_format.space_after = Pt(4)
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), LIGHT_FILL)
            p._p.get_or_add_pPr().append(shade)
        if node.type == "heading" and not quote:
            match = re.match(r"4\.(\d+) ", p.text)
            if match:
                nonlocal bookmark_count
                bookmark_count += 1
                start, end = (
                    OxmlElement("w:bookmarkStart"),
                    OxmlElement("w:bookmarkEnd"),
                )
                start.set(qn("w:id"), str(bookmark_count))
                start.set(qn("w:name"), f"report-rule-4-{match[1]}")
                end.set(qn("w:id"), str(bookmark_count))
                p._p.insert(1 if p._p.pPr is not None else 0, start)
                p._p.append(end)
        if node.type == "paragraph" and len(node.children or []) == 1:
            children = [t for t in node.children[0].token.children or [] if t.type != "text" or t.content]
            if (
                children
                and children[0].type == "strong_open"
                and children[-1].type == "strong_close"
            ):
                p.paragraph_format.keep_with_next = True
        return p

    def render(node, container, quote=False, depth=0):
        if node.type == "excerpt_table":
            # 表后同一行的文字或相邻表格也保留，不交给 HTML 吞掉 Markdown。
            parts = re.split(
                r"(<table\b[^>]*>.*?</table\s*>)", node.content, flags=re.I | re.S
            )
            for part in parts:
                if re.match(r"<table(?:\s|>)", part, re.I):
                    _render_excerpt_table(container, part, page_layout)
                elif part.strip():
                    for child in SyntaxTreeNode(parser.parse(part)).children:
                        render(child, container, quote, depth)
        elif node.type in {"paragraph", "heading", "code_block", "fence"}:
            paragraph(container, node, quote, depth=depth)
        elif node.type in {"ordered_list", "bullet_list"}:
            numbering = _numbering(
                document, node.type == "ordered_list", int(node.attrs.get("start", 1))
            )
            for item in node.children:
                numbered = False
                for child in item.children:
                    if child.type == "paragraph" and not numbered:
                        paragraph(container, child, quote, numbering, depth + 1)
                        numbered = True
                    else:
                        render(child, container, quote, depth + 1)
        elif node.type == "table":
            page_layout()
            rows = [row for section in node.children for row in section.children]
            columns = max(len(row.children) for row in rows)
            table = container.add_table(rows=0, cols=columns)
            table.style = "Table Grid"
            for row in rows:
                cells = table.add_row().cells
                for i, cell in enumerate(row.children):
                    for child in cell.children:
                        if child.type == "inline":
                            _inline(
                                cells[i].paragraphs[0],
                                child.token.children or [],
                                anchors=anchors,
                            )
            width = TABLE_WIDTH_DXA // columns
            set_table_geometry(
                table,
                [width] * (columns - 1) + [TABLE_WIDTH_DXA - width * (columns - 1)],
            )
            format_table_header(table.rows[0])
        else:
            for child in node.children or []:
                render(child, container, quote or node.type == "blockquote", depth)

    for node in nodes:
        render(node, document)
