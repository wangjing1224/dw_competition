"""将 XLSX 的每个 Sheet 解析为独立表格章节。"""

import hashlib
import re
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import StringIO
from html import escape
from pathlib import Path

from openpyxl import load_workbook

from app.core.config import get_settings
from app.models.documents import DocumentRef, ParsedBlock, ParsedDocument, ParsedSection
from app.parsers.assets import AssetStore
from app.parsers.common import assets_for, finish_document


def _display_value(cell) -> str:
    """保留公式；按常见 Excel 显示格式输出值，不模拟公式或猜测复杂格式。"""
    value = cell.value
    if value is None:
        return ""
    if cell.data_type == "f":
        return "[公式，未计算] " + str(value)
    fmt = cell.number_format
    if isinstance(value, (datetime, date, time)):
        # 仅处理明确的年月日/时分秒格式。区域、星期、条件和累计时长等
        # 复杂格式保持原值，避免为了排版改变日期含义。
        tokens = re.findall(r'"[^"]*"|[ymdhs]+|[-/: .年月日时分秒]', fmt, re.I)
        if "".join(tokens) == fmt and tokens:
            result = []
            fields = [t.lower() for t in tokens if t[0].lower() in "ymdhs"]
            index = 0
            for token in tokens:
                if token.startswith('"'):
                    result.append(token[1:-1])
                elif token.lower() in fields:
                    token = token.lower()
                    previous = fields[index - 1][0] if index else ""
                    following = fields[index + 1][0] if index + 1 < len(fields) else ""
                    index += 1
                    if token[0] == "m":
                        field = "minute" if previous == "h" or following == "s" else "month"
                    else:
                        field = {"y": "year", "d": "day", "h": "hour", "s": "second"}[token[0]]
                    if not hasattr(value, field) or len(token) > (4 if field == "year" else 2):
                        return str(value)
                    number = getattr(value, field)
                    if token == "yy":
                        number %= 100
                    result.append(str(number).zfill(len(token)))
                else:
                    result.append(token)
            return "".join(result)
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        number = re.fullmatch(r"(#,##)?(0+)(?:\.(0+))?(%)?", fmt)
        if number:
            comma, zeros, decimals, percent = number.groups()
            try:
                scaled = Decimal(str(value)) * (100 if percent else 1)
                rounded = scaled.quantize(Decimal(1).scaleb(-len(decimals or "")), rounding=ROUND_HALF_UP)
                rendered = format(rounded, ",f" if comma else "f")
            except InvalidOperation:
                return str(value)
            if not comma:
                sign = "-" if rendered.startswith("-") else ""
                length = len(zeros) + (len(decimals) + 1 if decimals else 0)
                rendered = sign + rendered.lstrip("-").zfill(length)
            return rendered + ("%" if percent else "")
    return str(value)


def _header_rows(sheet, covered) -> set[int]:
    """仅采信表格开头清晰的加粗列名带，不把首个非空行强认作表头。"""
    headers = set()
    preceding_dense = False
    for row in sheet.iter_rows():
        cells = [c for c in row if (c.row, c.column) not in covered and c.value is not None]
        if not cells:
            if headers:
                break
            preceding_dense = False
            continue
        is_header = (
            len(cells) >= 2
            and all(isinstance(c.value, str) and c.data_type != "f" and c.font.bold for c in cells)
        )
        if is_header:
            if preceding_dense and not headers:
                break
            headers.add(row[0].row)
        elif headers or any(not isinstance(c.value, str) or c.data_type == "f" for c in cells):
            break
        else:
            preceding_dense = len(cells) >= max(2, sheet.max_column * 0.65)
    return headers


def _column_widths(sheet) -> list[float]:
    widths = [sheet.sheet_format.defaultColWidth or 8.43] * sheet.max_column
    # 列分组可能只在首列登记 min/max，不能靠访问 column_dimensions 新建默认项。
    for dimension in sheet.column_dimensions.values():
        for column in range(dimension.min or 1, min(dimension.max or 1, sheet.max_column) + 1):
            if dimension.width and dimension.width > 0:
                widths[column - 1] = dimension.width
    return widths


def _section_id(file_id: str, sheet_name: str) -> str:
    digest = hashlib.sha1(f"{file_id}:{sheet_name}".encode()).hexdigest()[:20]
    return f"section-{digest}"


class XlsxParser:
    def supports(self, suffix: str) -> bool:
        return suffix.lower() == ".xlsx"

    def parse(
        self,
        document: DocumentRef,
        stored_path: Path,
        *,
        assets: AssetStore | None = None,
    ) -> ParsedDocument:
        # 普通模式才能读取合并区域与嵌入图片；公式保留原表达式，不能把无缓存值当空白。
        workbook = load_workbook(stored_path, data_only=False)
        assets = assets or assets_for(document)
        sections: list[ParsedSection] = []
        warnings = []
        try:
            for sheet in workbook.worksheets:
                blocks = []
                if (
                    sheet.max_row * sheet.max_column
                    > get_settings().parsing.max_sheet_cells
                ):
                    warnings.append(
                        f"工作表 {sheet.title} 超过单元格处理上限，未截断冒充完整表"
                    )
                    continue
                merges, covered = {}, set()
                for region in sheet.merged_cells.ranges:
                    merges[(region.min_row, region.min_col)] = (
                        region.max_row - region.min_row + 1,
                        region.max_col - region.min_col + 1,
                    )
                    covered.update(
                        (r, c)
                        for r in range(region.min_row, region.max_row + 1)
                        for c in range(region.min_col, region.max_col + 1)
                        if (r, c) != (region.min_row, region.min_col)
                    )
                headers = _header_rows(sheet, covered)
                rows = []
                for row in sheet.iter_rows():
                    cells = []
                    for cell in row:
                        if (cell.row, cell.column) in covered:
                            continue
                        rs, cs = merges.get((cell.row, cell.column), (1, 1))
                        value = _display_value(cell)
                        tag = "th" if cell.row in headers else "td"
                        # 原值/格式随解析快照保留，显示格式（如百分比、日期）不丢精度依据。
                        display = ""
                        if cell.value is not None and value != str(cell.value) and cell.data_type != "f":
                            display = (
                                f' data-raw-value="{escape(str(cell.value), quote=True)}"'
                                f' data-number-format="{escape(cell.number_format, quote=True)}"'
                            )
                        cells.append(
                            f'<{tag} rowspan="{rs}" colspan="{cs}"{display}>{escape(value)}</{tag}>'
                        )
                    rows.append(f'<tr data-source-row="{row[0].row}">' + "".join(cells) + "</tr>")
                if any(
                    cell.value is not None for row in sheet.iter_rows() for cell in row
                ):
                    blocks.append(
                        ParsedBlock(
                            kind="table",
                            content=(
                                '<table data-origin="worksheet">\n<colgroup>'
                                + "".join(f'<col data-width="{w:.2f}">' for w in _column_widths(sheet))
                                + "</colgroup>\n" + "\n".join(rows) + "\n</table>"
                            ),
                            source_block_index=0,
                        )
                    )
                for index, embedded in enumerate(sheet._images, start=1):
                    try:
                        anchor = getattr(embedded.anchor, "_from", None)
                        location = (
                            f"工作表 {sheet.title}，行 {anchor.row + 1} 列 {anchor.col + 1}"
                            if anchor
                            else f"工作表 {sheet.title}"
                        )
                        image = assets.save_image(
                            embedded._data(), description=location
                        )
                        blocks.append(
                            ParsedBlock(
                                kind="image", images=[image], source_block_index=index
                            )
                        )
                    except Exception as exc:
                        warnings.append(
                            f"工作表 {sheet.title} 图片未提取：{type(exc).__name__}"
                        )
                if not blocks:
                    continue
                sections.append(
                    ParsedSection(
                        section_id=_section_id(document.file_id, sheet.title),
                        title=f"Sheet: {sheet.title}",
                        blocks=blocks,
                    )
                )
        finally:
            workbook.close()
        return finish_document(document, stored_path, sections, warnings)


class XlsParser:
    """读取嵌入的 BIFF XLS 保存值及合并区域，不通过改扩展名伪装成 XLSX。"""

    def parse(self, document, stored_path, *, assets=None):
        import xlrd

        book = xlrd.open_workbook(
            file_contents=stored_path.read_bytes(), formatting_info=True,
            logfile=StringIO(), on_demand=True,
        )
        sections, warnings = [], []
        try:
            for sheet in book.sheets():
                if sheet.nrows * sheet.ncols > get_settings().parsing.max_sheet_cells:
                    warnings.append(f"工作表 {sheet.name} 超过单元格处理上限，未截断冒充完整表")
                    continue
                merges, covered = {}, set()
                for r0, r1, c0, c1 in sheet.merged_cells:
                    merges[r0, c0] = (r1 - r0, c1 - c0)
                    covered.update(
                        (r, c) for r in range(r0, r1) for c in range(c0, c1)
                        if (r, c) != (r0, c0)
                    )
                rows = []
                for r in range(sheet.nrows):
                    cells = []
                    for c in range(sheet.ncols):
                        if (r, c) in covered:
                            continue
                        cell = sheet.cell(r, c)
                        value = cell.value
                        if cell.ctype == xlrd.XL_CELL_DATE:
                            value = xlrd.xldate_as_datetime(value, book.datemode).isoformat(sep=" ")
                        elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                            value = "TRUE" if value else "FALSE"
                        elif cell.ctype == xlrd.XL_CELL_ERROR:
                            value = xlrd.error_text_from_code.get(value, "[单元格错误]")
                        elif cell.ctype == xlrd.XL_CELL_NUMBER and value.is_integer():
                            value = int(value)
                        rs, cs = merges.get((r, c), (1, 1))
                        content = escape(str(value)).replace("\n", "<br>")
                        cells.append(f'<td rowspan="{rs}" colspan="{cs}">{content}</td>')
                    rows.append("<tr>" + "".join(cells) + "</tr>")
                if rows:
                    sections.append(ParsedSection(
                        section_id=_section_id(document.file_id, sheet.name),
                        title=f"Sheet: {sheet.name}",
                        blocks=[ParsedBlock(
                            kind="table", content="<table>\n" + "\n".join(rows) + "\n</table>",
                            source_block_index=0,
                        )],
                    ))
        finally:
            book.release_resources()
        # xlrd 不读取图形/子附件，也不计算公式，不能把这些能力当作已完成。
        warnings.append("旧版 XLS 仅读取已保存单元格值及合并区域；公式未重算，图形与子附件未展开")
        return finish_document(document, stored_path, sections, warnings)
